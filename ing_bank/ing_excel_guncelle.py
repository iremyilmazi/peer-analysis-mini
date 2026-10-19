from __future__ import annotations

import re
import sys
import unicodedata
from copy import copy
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.views import Selection

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from entegre_excel_olustur import (
    extract_standard_rows_current_only,
    first_text,
    looks_like_financial_value,
    read_kap_excel_tables,
    table_title,
)

SHEET_NAME = "ing-bank"
TEMPLATE_CANDIDATES = ("akbank", "vakifbank", "turk-ekonomi-bankasi", "garanti-bankasi")
TOP_SCALE_MARKER = "ING_TOP_Q2_MILLION_TO_THOUSAND_V2"
PDF_SCALE_MARKER = "ING_PDF_Q2_MILLION_TO_THOUSAND_V2"


def normalize_label(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode().casefold()
    text = re.sub(r"[()\-+/]", " ", text)
    return " ".join(text.split())


def financial_number(value: object, multiply_q2: bool = False) -> object:
    if isinstance(value, (int, float)):
        number = value
    elif isinstance(value, str):
        text = value.strip().replace(" ", "")
        if text in {"", "-", "—", "–"}:
            return None
        negative = text.startswith("(") and text.endswith(")")
        if negative:
            text = text[1:-1]
        if not re.fullmatch(r"[+-]?\d[\d.,]*", text):
            return value
        number = int(re.sub(r"[.,]", "", text))
        if negative:
            number = -number
    else:
        return value
    return number * 1000 if multiply_q2 else number


def sequence_map(canonical: list[str], incoming: list[str]) -> dict[int, int]:
    left = [normalize_label(x) for x in canonical]
    right = [normalize_label(x) for x in incoming]
    matcher = SequenceMatcher(None, left, right, autojunk=False)
    mapping: dict[int, int] = {}
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            for offset in range(i2 - i1):
                mapping[i1 + offset] = j1 + offset
        elif tag == "replace" and (i2 - i1) == (j2 - j1):
            for offset in range(i2 - i1):
                if SequenceMatcher(None, left[i1 + offset], right[j1 + offset]).ratio() >= 0.72:
                    mapping[i1 + offset] = j1 + offset
    return mapping


def current_cumulative_column(table, report_period: str) -> int:
    """Select by reporting-period header, never by a parser-specific offset."""
    candidates = []
    for column in range(table.shape[1]):
        header = " ".join(str(value) for value in table.iloc[:2, column].tolist())
        normalized = normalize_label(header)
        current = re.search(r"\b(?:cari|current)\b", normalized)
        quarterly = re.search(r"\b(?:3|uc)\s*ayl|\b(?:quarter|three months?)\b", normalized)
        dates = re.findall(r"(?<!\d)\d{2}\.\d{2}\.\d{4}(?!\d)", header)
        if current and not quarterly and report_period in dates:
            candidates.append(column)
    if len(candidates) != 1:
        raise RuntimeError(
            f"ING cari kumulatif donem sutunu tekil bulunamadi: {report_period}; "
            f"adaylar={candidates}"
        )
    return candidates[0]


def income_rows(tables, report_period: str) -> list[dict[str, object]]:
    candidates = [item for item in tables
                  if len(item) > 5 and "kar veya zarar tablosu" in normalize_label(table_title(item))]
    if len(candidates) != 1:
        raise RuntimeError(f"ING gelir tablosu tekil bulunamadi: {len(candidates)}")
    table = candidates[0]
    value_column = current_cumulative_column(table, report_period)
    rows: list[dict[str, object]] = []
    for _, row in table.iloc[2:].iterrows():
        values = row.tolist()
        label = first_text(values[:value_column])
        current_cumulative = values[value_column]
        if label and looks_like_financial_value(current_cumulative):
            # An explicitly published dash in the income statement means zero;
            # a genuinely missing value is excluded and caught by coverage.
            if isinstance(current_cumulative, str) and current_cumulative.strip() in {"-", "–", "—"}:
                current_cumulative = 0
            rows.append({"label": label, "value": current_cumulative, "period": report_period})
    if not rows:
        raise RuntimeError(f"ING gelir tablosunda cari donem verisi bulunamadi: {report_period}")
    return rows


def copy_template_style(template, target) -> None:
    for row in range(1, 252):
        target.row_dimensions[row].height = template.row_dimensions[row].height
        for col in range(1, 80):
            source = template.cell(row, col)
            cell = target.cell(row, col)
            if source.has_style:
                cell._style = copy(source._style)
            cell.alignment = copy(source.alignment)
    for col_letter, dimension in template.column_dimensions.items():
        if 1 <= dimension.min <= 79:
            target.column_dimensions[col_letter].width = dimension.width
            target.column_dimensions[col_letter].hidden = dimension.hidden
    target.freeze_panes = template.freeze_panes
    if target.freeze_panes is None:
        target.sheet_view.selection = [Selection(activeCell="A1", sqref="A1")]
    target.sheet_view.showGridLines = template.sheet_view.showGridLines


def rebuild_ing_excel_pattern(workbook_path: Path, history_dir: Path) -> int:
    wb = load_workbook(workbook_path)
    if SHEET_NAME not in wb.sheetnames:
        raise RuntimeError(f"Sheet bulunamadı: {SHEET_NAME}")
    template_name = next((name for name in TEMPLATE_CANDIDATES if name in wb.sheetnames), None)
    if template_name is None:
        raise RuntimeError("Standart banka şablonu bulunamadı.")
    ws = wb[SHEET_NAME]
    template = wb[template_name]

    header_values = [[ws.cell(r, c).value for c in range(1, 80)] for r in (1, 2)]
    copy_template_style(template, ws)
    for row in range(3, 252):
        for col in range(1, 80):
            ws.cell(row, col).value = None
        ws.cell(row, 1).value = template.cell(row, 1).value
    for r_index, row in enumerate((1, 2)):
        for col in range(1, 80):
            ws.cell(row, col).value = header_values[r_index][col - 1]

    canonical_main_positions = [r for r in range(3, 188) if template.cell(r, 1).value]
    canonical_main = [str(template.cell(r, 1).value) for r in canonical_main_positions]
    canonical_income_positions = [r for r in range(190, 252) if template.cell(r, 1).value]
    canonical_income = [str(template.cell(r, 1).value) for r in canonical_income_positions]

    written_periods = 0
    for source in sorted(history_dir.rglob("*.xls*")):
        if not source.is_file() or source.name.startswith("~$"):
            continue
        tables = read_kap_excel_tables(source)
        report_period, main_rows = extract_standard_rows_current_only(tables)
        if not report_period or not main_rows:
            raise RuntimeError(f"ING standart tablo okunamadı: {source}")
        income = income_rows(tables, report_period)

        period_col = next(
            (col for col in range(2, 80, 3) if ws.cell(1, col).value == report_period),
            None,
        )
        if period_col is None:
            raise RuntimeError(f"Dönem sütunu bulunamadı: {report_period}")
        q2_million = report_period == "30.06.2026"

        main_labels = [str(row["label"]) for row in main_rows]
        for canonical_index, incoming_index in sequence_map(canonical_main, main_labels).items():
            row_no = canonical_main_positions[canonical_index]
            values = main_rows[incoming_index]["values"]
            for offset, breakdown in enumerate(("TL", "YP", "TOPLAM")):
                value = financial_number(values.get((report_period, breakdown)), multiply_q2=q2_million)
                cell = ws.cell(row_no, period_col + offset, value=value)
                if isinstance(value, (int, float)):
                    cell.number_format = "#,##0;[Red](#,##0);-"

        income_labels = [str(row["label"]) for row in income]
        income_mapping = sequence_map(canonical_income, income_labels)
        missing_income = [canonical_income[i] for i in range(len(canonical_income)) if i not in income_mapping]
        if missing_income:
            raise RuntimeError(
                f"ING gelir tablosu esleme eksik: {report_period}; "
                f"{len(missing_income)} satir; {missing_income[:5]}"
            )
        for canonical_index, incoming_index in income_mapping.items():
            row_no = canonical_income_positions[canonical_index]
            value = financial_number(income[incoming_index]["value"], multiply_q2=q2_million)
            cell = ws.cell(row_no, period_col + 2, value=value)
            if isinstance(value, (int, float)):
                cell.number_format = "#,##0;[Red](#,##0);-"
        written_periods += 1

    ws.cell(1, 200, TOP_SCALE_MARKER)
    if ws.cell(2, 200).value == "ING_Q2_MILLION_TO_THOUSAND_V1":
        ws.cell(2, 200, PDF_SCALE_MARKER)
    ws.column_dimensions["GR"].hidden = True
    wb.save(workbook_path)
    wb.close()
    return written_periods


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--history-dir", default=str(PROJECT_ROOT / "ing_bank" / "data" / "history"))
    args = parser.parse_args()
    print(rebuild_ing_excel_pattern(Path(args.workbook), Path(args.history_dir)))
