from __future__ import annotations

import argparse
import os
import shutil
import sys
from datetime import date
from pathlib import Path

COMPANY_NAME = "ing-bank"
PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CONFIG_PATH = PROJECT_ROOT / "kap_config.json"
DEFAULT_WORKBOOK_PATH = PROJECT_ROOT / "butunlesik_banka_verisi.xlsx"
DEFAULT_PDF_DIR = PROJECT_ROOT / "ing_bank" / "data" / "history"
DEFAULT_CACHE_DIR = PROJECT_ROOT / "ing_bank" / "data" / "pdf_cache"


def _financial_number(value: object) -> object:
    if not isinstance(value, str):
        return value
    text = value.strip().replace(" ", "")
    if text in {"", "-", "—", "–"}:
        return None
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1]
    import re
    if not re.fullmatch(r"[+-]?\d[\d.,]*", text):
        return value
    number = int(re.sub(r"[.,]", "", text))
    return -number if negative else number


def finalize_ing_sheet(workbook: Path, had_index: bool) -> None:
    from openpyxl import load_workbook
    from openpyxl.worksheet.views import Selection
    wb = load_workbook(workbook)
    if not had_index and "Index" in wb.sheetnames:
        del wb["Index"]
    ws = wb[COMPANY_NAME]
    # KAP Excel değerleri yerel ayırıcılarla metin gelebilir. Tüm dönemleri
    # hesaplanabilir sayılara çevir; PDF bloğu zaten sayısal yazılır.
    for row in range(3, min(252, ws.max_row) + 1):
        for col in range(2, min(79, ws.max_column) + 1):
            cell = ws.cell(row, col)
            converted = _financial_number(cell.value)
            if converted != cell.value:
                cell.value = converted
                if isinstance(converted, (int, float)):
                    cell.number_format = "#,##0;[Red](#,##0);-"
    top_marker = "ING_TOP_Q2_MILLION_TO_THOUSAND_V2"
    pdf_marker = "ING_PDF_Q2_MILLION_TO_THOUSAND_V2"
    top_marker_cell = ws.cell(1, 200)
    pdf_marker_cell = ws.cell(2, 200)
    legacy_scaled = pdf_marker_cell.value == "ING_Q2_MILLION_TO_THOUSAND_V1"
    if top_marker_cell.value != top_marker and not legacy_scaled:
        # 2026 Q2 KAP raporu milyon TL, önceki dönemler bin TL'dir.
        # Ana Excel alanlarında Q2 üçlüsü BY:BZ:CA'dır.
        for row in range(3, min(252, ws.max_row) + 1):
            for col in range(77, 80):
                cell = ws.cell(row, col)
                if isinstance(cell.value, (int, float)):
                    cell.value *= 1000
                    cell.number_format = "#,##0;[Red](#,##0);-"

        top_marker_cell.value = top_marker
    elif legacy_scaled:
        top_marker_cell.value = top_marker

    if pdf_marker_cell.value != pdf_marker and not legacy_scaled:
        # PDF bloklarında oran ve kur tabloları ölçeklenmez.
        current_title = ""
        for row in range(253, ws.max_row + 1):
            label = str(ws.cell(row, 1).value or "")
            next_label = str(ws.cell(row + 1, 1).value or "") if row < ws.max_row else ""
            if label and next_label.casefold() in {"kalem", "vade"}:
                current_title = label.casefold()
            skip_scale = "likidite karşılama" in current_title or "kur riskine" in current_title
            if skip_scale:
                continue
            for col in range(77, 80):
                cell = ws.cell(row, col)
                if isinstance(cell.value, (int, float)):
                    cell.value *= 1000
                    cell.number_format = "#,##0;[Red](#,##0);-"
        pdf_marker_cell.value = pdf_marker
    elif legacy_scaled:
        pdf_marker_cell.value = pdf_marker
    ws.column_dimensions["GR"].hidden = True
    # Bazı kaynak şablonlarda dondurulmuş bölme kaldırıldığı halde eski pane
    # seçimleri kalabiliyor. Excel bu çelişkili sheetView kaydını açılışta
    # onarmaya çalışır; bölme yoksa tek ve panesiz geçerli seçim bırak.
    if ws.freeze_panes is None:
        ws.sheet_view.selection = [Selection(activeCell="A1", sqref="A1")]
    wb.save(workbook)


def main() -> None:
    parser = argparse.ArgumentParser(description="ING Bank Excel ve PDF tarihçesini günceller.")
    parser.add_argument("--config", default=str(DEFAULT_CONFIG_PATH))
    parser.add_argument("--workbook", default=str(DEFAULT_WORKBOOK_PATH))
    parser.add_argument("--pdf-dir", default=str(DEFAULT_PDF_DIR))
    parser.add_argument("--cache-dir", default=str(DEFAULT_CACHE_DIR))
    parser.add_argument("--workers", type=int, default=4)
    parser.add_argument("--skip-download", action="store_true")
    parser.add_argument("--skip-excel", action="store_true")
    parser.add_argument("--skip-pdf", action="store_true")
    parser.add_argument("--all-history", action="store_true")
    args = parser.parse_args()
    os.chdir(PROJECT_ROOT)

    if str(PROJECT_ROOT) not in sys.path:
        sys.path.insert(0, str(PROJECT_ROOT))
    workbook = Path(args.workbook)
    pdf_dir = Path(args.pdf_dir)
    from openpyxl import load_workbook
    had_index = workbook.exists() and "Index" in load_workbook(workbook, read_only=True).sheetnames

    if not args.skip_download:
        from ing_bank.ing_kap_history_indir import run, run_incremental
        quarters = run(2020, pdf_dir, date.today()) if args.all_history else run_incremental(pdf_dir, date.today())
        print("KAP dönemleri:", ", ".join(quarters) if quarters else "yeni dönem yok")
    if not args.skip_excel:
        from entegre_excel_olustur import build_integrated_workbook, load_config, select_companies
        companies, source_dir, _ = load_config(Path(args.config))
        results = build_integrated_workbook(select_companies(companies, COMPANY_NAME), source_dir, workbook)
        for result in results:
            print(f"Excel alanları {result.status}: {result.sheet_name} ({result.rows_written} satır)")
        from ing_bank.ing_excel_guncelle import rebuild_ing_excel_pattern
        period_count = rebuild_ing_excel_pattern(workbook, pdf_dir)
        print(f"ING Excel dosya deseni OK: {period_count} dönem")
    if not args.skip_pdf:
        # PDF bloğu yeniden yazılacağı için yalnız PDF ölçek işaretini sıfırla.
        from openpyxl import load_workbook as _load_workbook
        _wb = _load_workbook(workbook)
        _wb[COMPANY_NAME].cell(2, 200).value = None
        _wb.save(workbook)
        from ing_bank.ing_dipnot_guncelle import append_ing_pdf_notes
        source, rows = append_ing_pdf_notes(workbook, pdf_dir, cache_dir=Path(args.cache_dir), workers=args.workers)
        print(f"PDF dipnotları OK: {source.name} ({rows} satır)")
    finalize_ing_sheet(workbook, had_index)
    print(f"ING güncellemesi tamamlandı: {workbook.resolve()}")


if __name__ == "__main__":
    main()
