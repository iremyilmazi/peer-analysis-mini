from __future__ import annotations

import argparse
import json
import re
from dataclasses import dataclass
from pathlib import Path

import pandas as pd
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


DEFAULT_CONFIG_PATH = "kap_config.json"
DEFAULT_OUTPUT_PATH = "butunlesik_banka_verisi.xlsx"
EXCEL_SUFFIXES = {".xls", ".xlsx", ".xlsm"}
PERIOD_DATES = ("31.03", "30.06", "30.09", "31.12")
PERIOD_BREAKDOWNS = ("TL", "YP", "TOPLAM")
DEFAULT_START_YEAR = 2020


@dataclass(frozen=True)
class CompanyConfig:
    name: str
    url: str
    selection_rule: str = "solo"
    source_dir: str | None = None


@dataclass(frozen=True)
class BuildResult:
    company_name: str
    sheet_name: str
    source_file: Path | None
    status: str
    rows_written: int = 0
    message: str = ""


def load_config(config_path: Path) -> tuple[list[CompanyConfig], Path, Path]:
    data = json.loads(config_path.read_text(encoding="utf-8"))
    companies = [
        CompanyConfig(
            name=item["name"],
            url=item["url"],
            selection_rule=item.get("selection_rule", "solo"),
            source_dir=item.get("source_dir"),
        )
        for item in data.get("companies", [])
    ]
    if not companies:
        raise RuntimeError("Config icinde companies listesi bos.")

    source_dir = Path(data.get("download_dir", "downloads/kap_finansal_raporlar"))
    output_path = Path(data.get("integrated_workbook_path", DEFAULT_OUTPUT_PATH))
    return companies, source_dir, output_path


def find_latest_excel(company_dir: Path) -> Path | None:
    if not company_dir.exists():
        return None

    candidates = [
        path
        for path in company_dir.rglob("*")
        if path.is_file()
        and path.suffix.lower() in EXCEL_SUFFIXES
        and not path.name.startswith("~$")
        and path.stat().st_size > 0
    ]
    if not candidates:
        return None
    return max(candidates, key=lambda path: path.stat().st_mtime)


def list_excel_files(company_dir: Path) -> list[Path]:
    if not company_dir.exists():
        return []
    return sorted(
        [
            path
            for path in company_dir.rglob("*")
            if path.is_file()
            and path.suffix.lower() in EXCEL_SUFFIXES
            and not path.name.startswith("~$")
            and path.stat().st_size > 0
        ],
        key=lambda path: (path.stat().st_mtime, str(path)),
    )


def source_dir_for_company(company: CompanyConfig, default_source_dir: Path) -> Path:
    if company.source_dir:
        return Path(company.source_dir)
    return default_source_dir / company.name


def read_kap_excel_tables(path: Path) -> list[pd.DataFrame]:
    first_bytes = path.read_bytes()[:64].lower()
    if first_bytes.startswith(b"<html") or first_bytes.startswith(b"<!doctype"):
        try:
            # KAP export files declare UTF-8 in their HTML meta tag.  Letting
            # lxml guess from a Windows path can decode Turkish characters as
            # U+FFFD (FA�Z, D��ER), which then prevents exact row matching and
            # silently drops P&L values during the incremental merge.
            tables = pd.read_html(path, header=None, encoding="utf-8")
        except ImportError:
            tables = read_html_tables_without_lxml(path)
    else:
        tables = [pd.read_excel(path, header=None)]

    cleaned: list[pd.DataFrame] = []
    for table in tables:
        table = table.dropna(how="all").dropna(axis=1, how="all")
        if not table.empty:
            cleaned.append(table.reset_index(drop=True))
    if not cleaned:
        raise RuntimeError(f"Okunabilir tablo bulunamadi: {path}")
    return cleaned


def read_html_tables_without_lxml(path: Path) -> list[pd.DataFrame]:
    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
    tables: list[pd.DataFrame] = []
    for html_table in soup.find_all("table"):
        rows: list[list[str]] = []
        max_columns = 0
        for html_row in html_table.find_all("tr"):
            values = [
                " ".join(cell.get_text(" ", strip=True).split())
                for cell in html_row.find_all(["th", "td"])
            ]
            if values:
                rows.append(values)
                max_columns = max(max_columns, len(values))
        if not rows:
            continue
        normalized_rows = [row + [None] * (max_columns - len(row)) for row in rows]
        tables.append(pd.DataFrame(normalized_rows))
    return tables


def parse_date_text(value: object) -> str | None:
    if value is None or pd.isna(value):
        return None
    matches = re.findall(r"(\d{2}\.\d{2}\.\d{4})", str(value))
    # Range headers such as 01.01.2026-30.06.2026 represent the report
    # ending period; using the first date silently drops P&L values.
    return matches[-1] if matches else None


def period_sort_key(period: str) -> tuple[int, int, int]:
    day, month, year = [int(part) for part in period.split(".")]
    return year, month, day


def standard_periods_until(max_period: str, start_year: int) -> list[str]:
    max_key = period_sort_key(max_period)
    periods = [
        f"{day_month}.{year}"
        for year in range(start_year, max_key[0] + 1)
        for day_month in PERIOD_DATES
        if period_sort_key(f"{day_month}.{year}") <= max_key
    ]
    return sorted(periods, key=period_sort_key)


def standard_periods_for(tables: list[pd.DataFrame], start_year: int = 2025) -> list[str]:
    years = {start_year}
    max_period: str | None = None
    for table in tables:
        for value in table.head(3).to_numpy().ravel():
            date_text = parse_date_text(value)
            if date_text:
                years.add(int(date_text[-4:]))
                if max_period is None or period_sort_key(date_text) > period_sort_key(max_period):
                    max_period = date_text

    return standard_periods_until(max_period or f"31.12.{max(years)}", start_year)


def standard_periods_from_dates(dates: set[str], start_year: int = DEFAULT_START_YEAR) -> list[str]:
    years = {start_year}
    years.update(int(date_text[-4:]) for date_text in dates)
    max_period = max(dates, key=period_sort_key) if dates else f"31.12.{max(years)}"
    return standard_periods_until(max_period, start_year)


def first_text(values: list[object]) -> str:
    for value in values:
        if value is None or pd.isna(value):
            continue
        text = " ".join(str(value).split())
        if text:
            return text
    return ""


def table_title(table: pd.DataFrame) -> str:
    for row_index in range(min(len(table), 5)):
        text = first_text(list(table.iloc[row_index, : min(table.shape[1], 3)]))
        if text and not parse_date_text(text) and text.upper() not in {"TP", "YP", "TOPLAM"}:
            return text
    return "Finansal Tablo"


def value_is_present(value: object) -> bool:
    return value is not None and not pd.isna(value) and str(value).strip() != ""


def compatible_financial_tables(tables: list[pd.DataFrame]) -> list[pd.DataFrame]:
    compatible: list[pd.DataFrame] = []
    for table in tables:
        # KAP HTML'i kullanilan parser/lxml surumune gore colspan'lari
        # genisleterek 13, birlestirerek 9 sutun uretebiliyor. Her iki yapida
        # da son alti sutun cari ve onceki donemin TP/YP/TOPLAM degerleridir.
        if table.shape[1] < 7 or len(table) < 3:
            continue
        header_values = table.iloc[0].tolist()
        header_dates = [parse_date_text(value) for value in header_values]
        if sum(1 for date_text in header_dates if date_text) < 2:
            continue
        is_income_statement = any("3 ayl" in str(value).casefold() for value in header_values)
        # Balance/off-balance tables need six value columns.  KAP's income
        # statement can legitimately be only eight columns with the four
        # period values at columns 3..6, depending on the installed
        # pandas/lxml versions and their rowspan/colspan expansion.
        if not is_income_statement and table.shape[1] < 9:
            continue
        if any(value_is_present(value) for value in table.iloc[:, -6:].to_numpy().ravel()):
            compatible.append(table)
    return compatible


def looks_like_financial_value(value: object) -> bool:
    if not value_is_present(value):
        return False
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return True
    text = str(value).strip().replace("\u00a0", "").replace(" ", "")
    if text in {"-", "–", "—"}:
        return True
    return bool(re.fullmatch(r"\(?[-+]?\d[\d.,]*\)?", text))


def income_statement_value_columns(table: pd.DataFrame) -> list[int]:
    """Locate the four income-statement period columns across parser versions."""
    scores: list[tuple[int, int]] = []
    for column in range(table.shape[1]):
        score = sum(looks_like_financial_value(value) for value in table.iloc[2:, column].tolist())
        if score:
            scores.append((score, column))
    # pandas/lxml can expand KAP's nested title table from 8 to 12 columns.
    # The four financial columns remain the four most numeric columns; sorting
    # their indices restores current cumulative, prior cumulative, current
    # quarter and prior quarter order.
    return sorted(column for _score, column in sorted(scores, reverse=True)[:4])


def extract_standard_rows(tables: list[pd.DataFrame]) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for table in compatible_financial_tables(tables):
        dates = [date_text for date_text in (parse_date_text(value) for value in table.iloc[0].tolist()) if date_text]
        if len(dates) < 2:
            continue

        current_period = dates[0]
        previous_period = dates[1]
        for _, row in table.iloc[2:].iterrows():
            row_values = row.tolist()
            value_block = row_values[-6:]
            if not any(value_is_present(value) for value in value_block):
                continue

            label = first_text(row_values[:-6])
            if not label:
                continue

            values = {
                (current_period, "TL"): value_block[0],
                (current_period, "YP"): value_block[1],
                (current_period, "TOPLAM"): value_block[2],
                (previous_period, "TL"): value_block[3],
                (previous_period, "YP"): value_block[4],
                (previous_period, "TOPLAM"): value_block[5],
            }
            rows.append({"label": label, "values": values})

    return rows


def extract_standard_rows_current_only(tables: list[pd.DataFrame]) -> tuple[str | None, list[dict[str, object]]]:
    rows: list[dict[str, object]] = []
    report_period: str | None = None
    for table in compatible_financial_tables(tables):
        dates = [date_text for date_text in (parse_date_text(value) for value in table.iloc[0].tolist()) if date_text]
        if not dates:
            continue

        unique_dates = list(dict.fromkeys(dates))
        current_period = unique_dates[0]
        report_period = report_period or current_period
        is_income_statement = any("3 ayl" in str(value).casefold() for value in table.iloc[0].tolist())
        income_value_columns = income_statement_value_columns(table) if is_income_statement else []
        for _, row in table.iloc[2:].iterrows():
            row_values = row.tolist()
            if is_income_statement:
                # Income-statement tables carry current cumulative, prior
                # cumulative, current quarter and prior quarter rather than
                # TL/YP/TOPLAM.  Store the current cumulative amount in the
                # integrated TOPLAM column.
                if len(income_value_columns) != 4:
                    continue
                values = [row_values[column] for column in income_value_columns]
                if not any(looks_like_financial_value(value) for value in values):
                    continue
                first_value_column = income_value_columns[0]
                label = first_text(row_values[:first_value_column])
                if not label:
                    continue
                rows.append(
                    {
                        "label": label,
                        "values": {(current_period, "TOPLAM"): values[0]},
                    }
                )
                continue
            value_block = row_values[-6:-3]
            if not any(value_is_present(value) for value in value_block):
                continue

            label = first_text(row_values[:-6])
            if not label:
                continue

            rows.append(
                {
                    "label": label,
                    "values": {
                        (current_period, "TL"): value_block[0],
                        (current_period, "YP"): value_block[1],
                        (current_period, "TOPLAM"): value_block[2],
                    },
                }
            )

    return report_period, rows


def sheet_name_for(company_name: str, used_names: set[str]) -> str:
    cleaned = re.sub(r"[\[\]:*?/\\]", "_", company_name).strip()
    cleaned = cleaned[:31] or "Sheet"
    candidate = cleaned
    index = 2
    while candidate in used_names:
        suffix = f"_{index}"
        candidate = f"{cleaned[:31 - len(suffix)]}{suffix}"
        index += 1
    used_names.add(candidate)
    return candidate


def load_or_create_workbook(output_path: Path) -> Workbook:
    if output_path.exists():
        return load_workbook(output_path)
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)
    return workbook


def reset_sheet(workbook: Workbook, sheet_name: str):
    if sheet_name in workbook.sheetnames:
        index = workbook.sheetnames.index(sheet_name)
        workbook.remove(workbook[sheet_name])
        return workbook.create_sheet(sheet_name, index)
    return workbook.create_sheet(sheet_name)


def write_company_sheet(
    workbook: Workbook,
    company: CompanyConfig,
    sheet_name: str,
    source_file: Path,
    tables: list[pd.DataFrame],
) -> int:
    sheet = reset_sheet(workbook, sheet_name)
    periods = standard_periods_for(tables)
    standard_rows = extract_standard_rows(tables)
    if not standard_rows:
        raise RuntimeError(f"Standart TL/YP/TOPLAM tablo yapisi bulunamadi: {source_file}")

    write_standard_header(sheet, periods)

    current_row = 3
    for item in standard_rows:
        sheet.cell(row=current_row, column=1, value=item["label"])
        values = item["values"]
        for period_index, period in enumerate(periods):
            for breakdown_index, breakdown in enumerate(PERIOD_BREAKDOWNS):
                column_index = 2 + period_index * len(PERIOD_BREAKDOWNS) + breakdown_index
                sheet.cell(row=current_row, column=column_index, value=values.get((period, breakdown)))
        current_row += 1
        if item["label"] == "VARLIKLAR TOPLAMI":
            current_row += 2

    format_company_sheet(sheet)
    return len(standard_rows)


def write_company_history_sheet(
    workbook: Workbook,
    company: CompanyConfig,
    sheet_name: str,
    source_files: list[Path],
) -> tuple[int, list[str]]:
    sheet = reset_sheet(workbook, sheet_name)
    # Aynı görünen satır adı bir KAP tablosunda farklı hiyerarşiler altında
    # birden fazla kez bulunabilir (örn. "Devlet Borçlanma Senetleri",
    # "Diğer", "Krediler"). Yalnızca etiketi anahtar yapmak son örneğin
    # önceki örneğin değerlerini ezmesine neden olur. Etiket + dosya içindeki
    # tekrar sırası kalıcı satır anahtarıdır.
    row_order: list[tuple[str, int]] = []
    values_by_row_key: dict[tuple[str, int], dict[tuple[str, str], object]] = {}
    periods_seen: set[str] = set()
    warnings: list[str] = []

    for source_file in source_files:
        try:
            tables = read_kap_excel_tables(source_file)
            report_period, rows = extract_standard_rows_current_only(tables)
            if not report_period or not rows:
                raise RuntimeError("Standart TL/YP/TOPLAM cari donem tablo yapisi bulunamadi.")
            periods_seen.add(report_period)
            occurrence_by_label: dict[str, int] = {}
            for item in rows:
                label = str(item["label"])
                occurrence = occurrence_by_label.get(label, 0) + 1
                occurrence_by_label[label] = occurrence
                row_key = (label, occurrence)
                if row_key not in values_by_row_key:
                    values_by_row_key[row_key] = {}
                    row_order.append(row_key)
                values_by_row_key[row_key].update(item["values"])
        except Exception as exc:
            message = f"WARN: {company.name} Excel okunamadi: {source_file} -> {type(exc).__name__}: {exc}"
            print(message)
            warnings.append(message)

    if not row_order:
        raise RuntimeError("History icinde okunabilir Garanti Excel verisi bulunamadi.")

    periods = standard_periods_from_dates(periods_seen)
    write_standard_header(sheet, periods)

    current_row = 3
    for label, occurrence in row_order:
        sheet.cell(row=current_row, column=1, value=label)
        values = values_by_row_key[(label, occurrence)]
        for period_index, period in enumerate(periods):
            for breakdown_index, breakdown in enumerate(PERIOD_BREAKDOWNS):
                column_index = 2 + period_index * len(PERIOD_BREAKDOWNS) + breakdown_index
                sheet.cell(row=current_row, column=column_index, value=values.get((period, breakdown)))
        current_row += 1
        if label == "VARLIKLAR TOPLAMI":
            current_row += 2

    format_company_sheet(sheet)
    return len(row_order), warnings


def write_standard_header(sheet, periods: list[str]) -> None:
    sheet.cell(row=1, column=1, value="Kalem")
    sheet.cell(row=2, column=1, value="Kalem")
    for period_index, period in enumerate(periods):
        start_column = 2 + period_index * len(PERIOD_BREAKDOWNS)
        for breakdown_index, breakdown in enumerate(PERIOD_BREAKDOWNS):
            column_index = start_column + breakdown_index
            sheet.cell(row=1, column=column_index, value=period)
            sheet.cell(row=2, column=column_index, value=breakdown)


def format_company_sheet(sheet) -> None:
    sheet.freeze_panes = "B3"
    sheet.sheet_view.showGridLines = False

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    for row in sheet.iter_rows(min_row=1, max_row=2):
        for cell in row:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in sheet.iter_rows():
        for cell in row:
            if cell.row > 2:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    for column_cells in sheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = 0
        for cell in column_cells[:80]:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
        if column_letter == "A":
            sheet.column_dimensions[column_letter].width = min(max(max_length + 2, 28), 70)
        else:
            sheet.column_dimensions[column_letter].width = 14


def write_index_sheet(workbook: Workbook, results: list[BuildResult], output_path: Path) -> None:
    if "Index" in workbook.sheetnames:
        workbook.remove(workbook["Index"])
    sheet = workbook.create_sheet("Index", 0)
    headers = ["Banka", "Sheet", "Durum", "Kaynak dosya", "Satir", "Not"]
    sheet.append(headers)
    for result in results:
        sheet.append(
            [
                result.company_name,
                result.sheet_name,
                result.status,
                str(result.source_file) if result.source_file else "",
                result.rows_written,
                result.message,
            ]
        )

    sheet["A1"] = "Banka"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False

    for column_index, column_cells in enumerate(sheet.columns, start=1):
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[get_column_letter(column_index)].width = min(max(max_length + 2, 12), 60)

    sheet["H1"] = "Workbook"
    sheet["I1"] = str(output_path)
    sheet["H1"].font = Font(bold=True)


def build_integrated_workbook(
    companies: list[CompanyConfig],
    source_dir: Path,
    output_path: Path,
    *,
    update_index: bool = True,
) -> list[BuildResult]:
    workbook = load_or_create_workbook(output_path)
    existing_company_sheets = {name for name in workbook.sheetnames if name != "Index"}
    used_sheet_names = set(workbook.sheetnames)
    sheet_name_map = {name: name for name in existing_company_sheets}

    results: list[BuildResult] = []
    for company in companies:
        sheet_name = sheet_name_map.get(company.name) or sheet_name_for(company.name, used_sheet_names)
        company_source_dir = source_dir_for_company(company, source_dir)
        source_files = list_excel_files(company_source_dir)
        if company.source_dir and source_files:
            try:
                rows_written, warnings = write_company_history_sheet(workbook, company, sheet_name, source_files)
                results.append(
                    BuildResult(
                        company_name=company.name,
                        sheet_name=sheet_name,
                        source_file=source_files[-1],
                        status="OK" if not warnings else "WARN",
                        rows_written=rows_written,
                        message=f"{len(source_files)} history Excel islendi; {len(warnings)} uyari.",
                    )
                )
            except Exception as exc:
                results.append(
                    BuildResult(
                        company_name=company.name,
                        sheet_name=sheet_name,
                        source_file=source_files[-1] if source_files else None,
                        status="ERROR",
                        message=f"{type(exc).__name__}: {exc}",
                    )
                )
            continue

        source_file = find_latest_excel(company_source_dir)
        if source_file is None:
            results.append(
                BuildResult(
                    company_name=company.name,
                    sheet_name=sheet_name,
                    source_file=None,
                    status="SKIP",
                    message="Kaynak Excel bulunamadi.",
                )
            )
            continue

        try:
            tables = read_kap_excel_tables(source_file)
            rows_written = write_company_sheet(workbook, company, sheet_name, source_file, tables)
            results.append(
                BuildResult(
                    company_name=company.name,
                    sheet_name=sheet_name,
                    source_file=source_file,
                    status="OK",
                    rows_written=rows_written,
                )
            )
        except Exception as exc:
            results.append(
                BuildResult(
                    company_name=company.name,
                    sheet_name=sheet_name,
                    source_file=source_file,
                    status="ERROR",
                    message=f"{type(exc).__name__}: {exc}",
                )
            )

    # Tek banka guncellemelerinde mevcut Index sayfasini yeniden kurmak hem
    # diger banka satirlarini siler hem de gecici dosya yolunu workbook'a
    # yazabilir. Otomatik guncelleme akisi bu nedenle Index'i aynen korur.
    if update_index:
        write_index_sheet(workbook, results, output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return results


def select_companies(companies: list[CompanyConfig], company_name: str | None) -> list[CompanyConfig]:
    if not company_name:
        return companies
    selected = [company for company in companies if company.name == company_name]
    if not selected:
        raise RuntimeError(f"Config icinde company bulunamadi: {company_name}")
    return selected


def main() -> None:
    parser = argparse.ArgumentParser(description="KAP Excel ciktilarini banka bazli sheetlere sahip butunlesik workbook'a aktarir.")
    parser.add_argument("--config", default=DEFAULT_CONFIG_PATH, help="KAP ayar dosyasi.")
    parser.add_argument("--company", help="Sadece belirtilen config name degerine sahip bankayi workbook'a aktarir.")
    parser.add_argument("--source-dir", help="KAP Excel dosyalarinin bulundugu ana klasor.")
    parser.add_argument("--output", help="Olusturulacak/guncellenecek butunlesik Excel dosyasi.")
    args = parser.parse_args()

    companies, source_dir, output_path = load_config(Path(args.config))
    if args.source_dir:
        source_dir = Path(args.source_dir)
    if args.output:
        output_path = Path(args.output)

    selected_companies = select_companies(companies, args.company)
    results = build_integrated_workbook(selected_companies, source_dir, output_path)

    for result in results:
        source = result.source_file.name if result.source_file else "-"
        detail = f" ({result.message})" if result.message else ""
        print(f"{result.status}: {result.company_name} -> {result.sheet_name} [{source}] {result.rows_written} satir{detail}")
    print(f"Butunlesik Excel: {output_path.resolve()}")


if __name__ == "__main__":
    main()
