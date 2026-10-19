from __future__ import annotations

import argparse
import re
import unicodedata
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


TABLE_TITLE = (
    "Birinci ve ikinci aşama krediler, diğer alacaklar ile yeniden "
    "yapılandırılan ya da yeni bir itfa planına bağlanan krediler ve "
    "bunlar için ayrılan karşılıklara ilişkin bilgiler"
)
OUTPUT_FILE = "birinci_ikinci_asama_krediler.xlsx"
SHEET_NAME = "garanti"


def normalize_text(value: str) -> str:
    return " ".join(value.split())


def title_in_page(target: str, text: str) -> bool:
    def key(value: str) -> str:
        value = unicodedata.normalize("NFKD", normalize_text(value).casefold().replace("ı", "i"))
        value = "".join(ch for ch in value if not unicodedata.combining(ch))
        return " ".join(re.sub(r"[^a-z0-9]+", " ", value).split())
    target_tokens = key(target).split(); page_tokens = set(key(text).split())
    return key(target) in key(text) or (
        len(target_tokens) >= 3 and all(t in page_tokens for t in target_tokens[:3])
        and sum(t in page_tokens for t in target_tokens) / len(target_tokens) >= 0.50
    )


def find_default_pdf() -> Path:
    pdfs = sorted(Path(".").glob("*.pdf"))
    if not pdfs:
        raise FileNotFoundError("Bu klasörde PDF bulunamadı.")
    if len(pdfs) > 1:
        raise ValueError(
            "Klasörde birden fazla PDF var. Lütfen PDF yolunu argüman olarak verin."
        )
    return pdfs[0]


def extract_report_date(pdf_path: Path) -> str:
    match = re.search(r"(\d{2})\s+(\d{2})\s+(\d{4})", pdf_path.stem)
    if not match:
        return ""
    day, month, year = match.groups()
    return f"{day}.{month}.{year}"


def clean_cell(value: str | None) -> str:
    if value is None:
        return ""
    return normalize_text(value.replace("\n", " "))


def to_excel_number(value):
    if not isinstance(value, str):
        return value
    cleaned = normalize_text(value)
    if re.fullmatch(r"\d+(?:,\d{3})*", cleaned):
        return int(cleaned.replace(",", ""))
    return value


def extract_current_period_table(pdf_path: Path) -> list[list[str]]:
    target = normalize_text(TABLE_TITLE)

    with pdfplumber.open(str(pdf_path)) as pdf:
        for page in pdf.pages:
            text = page.extract_text(x_tolerance=1, y_tolerance=3) or ""
            if not title_in_page(target, text):
                continue

            for table in page.extract_tables():
                if not table or not table[0]:
                    continue
                if clean_cell(table[0][0]) != "Cari Dönem":
                    continue
                labels = {clean_cell(row[0]) for row in table[3:] if row and row[0]}
                if {"İhtisas Dışı Krediler", "Toplam"}.issubset(labels):
                    return [[clean_cell(cell) for cell in row] for row in table]

    raise RuntimeError(f"Hedef cari dönem tablosu PDF içinde bulunamadı: {TABLE_TITLE}")


def write_excel(table: list[list[str]], output_path: Path, report_date: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    header_fill = PatternFill("solid", fgColor="D9EAF7")

    for row_idx, row in enumerate(table, start=1):
        output_row = list(row)
        if row_idx == 1:
            output_row[0] = report_date

        for col_idx, value in enumerate(output_row, start=1):
            cell = ws.cell(row_idx, col_idx, to_excel_number(value))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if row_idx <= 3:
                cell.font = Font(bold=True)
                cell.fill = header_fill
                cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
            elif col_idx > 1:
                cell.alignment = Alignment(horizontal="right")
                if isinstance(cell.value, int):
                    cell.number_format = "#,##0"

    widths = {
        "A": 30,
        "B": 22,
        "C": 26,
        "D": 26,
        "E": 20,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    for row_idx in range(1, 4):
        ws.row_dimensions[row_idx].height = 34

    for col_idx in range(1, 6):
        col = get_column_letter(col_idx)
        for cell in ws[col]:
            if isinstance(cell.value, int):
                cell.number_format = "#,##0"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "Birinci ve ikinci aşama krediler tablosunun cari dönem bölümünü "
            "Excel'e aktarır."
        )
    )
    parser.add_argument("pdf", nargs="?", type=Path, help="Kaynak PDF yolu")
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=Path(OUTPUT_FILE),
        help="Oluşturulacak Excel dosyası",
    )
    args = parser.parse_args()

    pdf_path = args.pdf or find_default_pdf()
    table = extract_current_period_table(pdf_path)
    report_date = extract_report_date(pdf_path)
    write_excel(table, args.output, report_date)
    print(f"{len(table) - 3} cari dönem satırı yazıldı: {args.output}")


if __name__ == "__main__":
    main()
