from __future__ import annotations

import argparse
import concurrent.futures
import importlib.util
import json
import multiprocessing
import re
import sys
import tempfile
import hashlib
import time
from difflib import SequenceMatcher
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


COMPANY_NAME = "garanti-bankasi"
DEFAULT_WORKBOOK_PATH = "butunlesik_banka_verisi.xlsx"
DEFAULT_PDF_DIR = "garanti_bankasi/data/history"
DEFAULT_CACHE_DIR = "garanti_bankasi/data/pdf_cache"
DEFAULT_WORKERS = 4
PEERPDF_SCRIPT_PATH = Path(__file__).resolve().parent / "peerpdf_scripts" / "garanti_dipnot.py"
BLOCK_TITLE = "PDF Dipnot Alanlari"
PDF_BLOCK_START_ROW = 253
MARKER_COLUMN = 200
HEADER_FILL = "548235"
SOURCE_FILL = "E2F0D9"
MAX_CELL_TEXT_LENGTH = 32000
CACHE_VERSION = 11
EXTRACTOR_TIMEOUT_SECONDS = 90
CURRENT_PDF_EXTRACTOR_WORKERS = 3
PDF_PAGE_TEXT_CACHE_VERSION = 1
_PDF_TEXT_MEMORY_CACHE: dict[tuple[str, int, int], list[str]] = {}
STAGE_PROVISION_TITLE = (
    "Birinci ve ikinci asama krediler, diger alacaklar ile yeniden yapilandirilan "
    "ya da yeni bir itfa planina baglanan krediler ve bunlar icin ayrilan "
    "karsiliklara iliskin bilgiler"
)
STAGE_PROVISION_TARGET_LABELS = {
    "12 Aylik Beklenen Zarar Karsiligi",
    "Kredi Riskinde Onemli Artis",
}

TARGET_SECTIONS = [
    (1, 0, "consumer", "Tüketici kredileri, bireysel kredi kartları, personel kredileri ve personel kredi kartlarına ilişkin bilgiler", ["Kalem", "Kısa Vadeli", "Orta ve Uzun Vadeli", "Toplam"]),
    (2, 0, "installment", "Taksitli ticari krediler ve kurumsal kredi kartlarına ilişkin bilgiler", ["Kalem", "Kısa Vadeli", "Orta ve Uzun Vadeli", "Toplam"]),
    (3, 0, "deposit_maturity", "Mevduatın/toplanan fonların vade yapısına ilişkin bilgiler", ["Kalem", "Vadesiz", "Vadeli", "Toplam"]),
    (3, 1, "deposit_maturity_total", "Mevduatın vade yapısına ilişkin toplam bilgileri", ["Vade", "", "", "Toplam"]),
    (4, 0, "deposit_interest", "Mevduata ödenen faizin vade yapısına göre gösterimi", ["Kalem", "Vadesiz", "Vadeli", "Toplam"]),
    (4, 1, "deposit_interest_total", "Mevduata ödenen faizin vade yapısına göre gösterimi - toplam vade kırılımı", ["Vade", "", "", "Toplam"]),
    (5, 0, "restructured", "Donuk alacaklardan yeniden yapılandırılan kredilere ilişkin bilgiler", ["Kalem", "III. Grup", "IV. Grup", "V. Grup"]),
    (6, 0, "npl_movement", "Toplam donuk alacak hareketlerine ilişkin bilgiler", ["Kalem", "III. Grup", "IV. Grup", "V. Grup"]),
    (7, 0, "banks_interest", "Bankalardan alınan faiz gelirlerine ilişkin bilgiler", ["Kalem", "TP", "YP"]),
    (8, 0, "securities_interest", "Menkul değerlerden alınan faiz gelirlerine ilişkin bilgiler", ["Kalem", "TP", "YP"]),
    (9, 0, "equity", "Özkaynak kalemlerine ilişkin bilgiler", ["Kalem", "", "", "Toplam"]),
    (10, 0, "expected_loss", "Beklenen zarar karşılıkları ve diğer karşılık giderleri", ["Kalem", "", "", "Toplam"]),
    (11, 0, "stage12", "Birinci ve ikinci aşama krediler ve bunlar için ayrılan karşılıklara ilişkin bilgiler", ["Kalem", "Standart Nitelikli Krediler", "Yakın İzlemedeki Krediler"]),
    (13, 0, "liquidity", "Likidite karşılama oranı", ["Kalem", "TP+YP", "YP"]),
    (14, 0, "stage3", "Kredilere ilişkin olarak ayrılan 3. aşama karşılıklar", ["Kalem", "", "", "Toplam"]),
    (19, 0, "fx_risk", "Kur riskine ilişkin açıklamalar", ["Kalem", "Toplam"]),
    (19, 1, "credit_summary", "Birinci ve ikinci grup krediler, diğer alacaklar ile yeniden yapılandırılan ya da yeni bir itfa planına bağlanan krediler ve bunlar için ayrılan karşılıklara ilişkin bilgiler", ["Kalem", "TL", "YP"]),
    (20, 0, "other_income", "Diğer faaliyet gelirlerine ilişkin açıklamalar", ["Kalem", "", "", "Toplam"]),
    (21, 0, "other_expense", "Diğer faaliyet giderlerine ilişkin bilgiler", ["Kalem", "", "", "Toplam"]),
    (16, 0, "cash", "Nakit değerler ve TCMB’ye ilişkin bilgiler", ["Kalem", "TP", "YP"]),
    (17, 0, "used_loans", "Kullanılan kredilere verilen faizlere ilişkin bilgiler", ["Kalem", "TP", "YP"]),
    (18, 0, "loan_interest", "Kredilerden alınan faiz gelirlerine ilişkin bilgiler", ["Kalem", "TP", "YP"]),
]


@dataclass(frozen=True)
class PdfNote:
    note_no: str
    title: str
    text: str


def find_latest_pdf(pdf_dir: Path) -> Path:
    candidates = [
        path
        for path in pdf_dir.rglob("*.pdf")
        if path.is_file() and not path.name.startswith("~$") and path.stat().st_size > 0
    ]
    if not candidates:
        raise FileNotFoundError(f"Garanti PDF dosyasi bulunamadi: {pdf_dir}")
    return max(candidates, key=lambda path: path.stat().st_mtime)


def list_pdf_files(pdf_dir: Path) -> list[Path]:
    if not pdf_dir.exists():
        return []
    return sorted(
        [
            path
            for path in pdf_dir.rglob("*.pdf")
            if path.is_file() and not path.name.startswith("~$") and path.stat().st_size > 0
        ],
        key=lambda path: period_sort_key(extract_report_date(path) or "31.12.9999"),
    )


def extract_report_date(pdf_path: Path) -> str:
    match = re.search(r"(\d{2})\s+(\d{2})\s+(\d{4})", pdf_path.stem)
    if not match:
        match = re.search(r"(\d{2})[._-](\d{2})[._-](\d{4})", pdf_path.stem)
    if not match:
        return ""
    day, month, year = match.groups()
    return f"{day}.{month}.{year}"


def period_sort_key(period: str) -> tuple[int, int, int]:
    day, month, year = [int(part) for part in period.split(".")]
    return year, month, day


def previous_period_for(period: str) -> str:
    day, month, year = period.split(".")
    if day == "31" and month == "03":
        return f"31.12.{int(year) - 1}"
    if day == "30" and month == "06":
        return f"31.03.{year}"
    if day == "30" and month == "09":
        return f"30.06.{year}"
    if day == "31" and month == "12":
        return f"30.09.{year}"
    raise RuntimeError(f"Onceki donem hesaplanamadi: {period}")


def period_start_column(sheet, period: str) -> int:
    if not period:
        raise RuntimeError("PDF rapor donemi dosya adindan okunamadi.")
    for column_index in range(2, sheet.max_column + 1):
        if sheet.cell(row=1, column=column_index).value == period:
            return column_index
    raise RuntimeError(f"Workbook icinde PDF donemi icin kolon bulunamadi: {period}")


def normalize_key(value: str) -> str:
    replacements = str.maketrans(
        {
            "ç": "c",
            "Ç": "C",
            "ğ": "g",
            "Ğ": "G",
            "ı": "i",
            "I": "I",
            "İ": "I",
            "ö": "o",
            "Ö": "O",
            "ş": "s",
            "Ş": "S",
            "ü": "u",
            "Ü": "U",
        }
    )
    return " ".join(value.translate(replacements).split()).casefold()


def _pdf_memory_key(pdf_path: Path) -> tuple[str, int, int]:
    stat = pdf_path.stat()
    return str(pdf_path.resolve()), stat.st_size, stat.st_mtime_ns


def _file_sha256(pdf_path: Path) -> str:
    digest = hashlib.sha256()
    with pdf_path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _page_text_cache_path(source_pdf_path: Path, payload_sha256: str) -> Path:
    cache_dir = source_pdf_path.parent / "cache"
    return cache_dir / f"{source_pdf_path.stem}.{payload_sha256[:16]}.page_text.json"


def load_or_build_pdf_page_texts(
    source_pdf_path: Path,
    payload_pdf_path: Path | None = None,
) -> list[str]:
    """Extract every page once and persist the result beside the history PDF."""
    add_bundled_pdf_dependencies()
    import pdfplumber

    payload_pdf_path = payload_pdf_path or materialize_pdf_payload(source_pdf_path)
    memory_key = _pdf_memory_key(payload_pdf_path)
    cached_memory = _PDF_TEXT_MEMORY_CACHE.get(memory_key)
    if cached_memory is not None:
        return cached_memory

    payload_sha256 = _file_sha256(payload_pdf_path)
    cache_path = _page_text_cache_path(source_pdf_path, payload_sha256)
    try:
        payload = json.loads(cache_path.read_text(encoding="utf-8"))
        if (
            payload.get("version") == PDF_PAGE_TEXT_CACHE_VERSION
            and payload.get("payload_sha256") == payload_sha256
            and isinstance(payload.get("pages"), list)
        ):
            page_texts = [str(text or "") for text in payload["pages"]]
            _PDF_TEXT_MEMORY_CACHE[memory_key] = page_texts
            print(
                f"PDF sayfa metni onbellekten yuklendi: {len(page_texts)} sayfa",
                flush=True,
            )
            return page_texts
    except (OSError, ValueError, TypeError):
        pass

    started = time.monotonic()
    page_texts: list[str] = []
    print("PDF sayfa metni tek geciste indeksleniyor...", flush=True)
    with pdfplumber.open(str(payload_pdf_path)) as pdf:
        total_pages = len(pdf.pages)
        for page_index, page in enumerate(pdf.pages, start=1):
            page_texts.append(
                page.extract_text(x_tolerance=1, y_tolerance=3) or ""
            )
            if page_index % 25 == 0 or page_index == total_pages:
                print(
                    f"PDF sayfa indeksi: {page_index}/{total_pages}",
                    flush=True,
                )

    cache_path.parent.mkdir(parents=True, exist_ok=True)
    cache_payload = {
        "version": PDF_PAGE_TEXT_CACHE_VERSION,
        "payload_sha256": payload_sha256,
        "pages": page_texts,
    }
    temporary_path = cache_path.with_suffix(cache_path.suffix + ".tmp")
    temporary_path.write_text(
        json.dumps(cache_payload, ensure_ascii=False),
        encoding="utf-8",
    )
    temporary_path.replace(cache_path)
    _PDF_TEXT_MEMORY_CACHE[memory_key] = page_texts
    print(
        f"PDF sayfa indeksi tamamlandi: {len(page_texts)} sayfa, "
        f"{time.monotonic() - started:.1f} sn",
        flush=True,
    )
    return page_texts


def read_pdf_text(pdf_path: Path) -> str:
    add_bundled_pdf_dependencies()
    page_texts = load_or_build_pdf_page_texts(pdf_path)
    combined = "\n".join(page_texts)
    if not combined.strip():
        raise RuntimeError(f"PDF metni okunamadi: {pdf_path}")
    return combined


def add_bundled_pdf_dependencies() -> None:
    """Compatibility hook; install dependencies from requirements.txt."""
    return None


def load_peerpdf_garanti_module():
    add_bundled_pdf_dependencies()
    if not PEERPDF_SCRIPT_PATH.exists():
        raise FileNotFoundError(f"PeerPDF Garanti parser bulunamadi: {PEERPDF_SCRIPT_PATH}")

    module_dir = str(PEERPDF_SCRIPT_PATH.parent)
    if module_dir not in sys.path:
        sys.path.insert(0, module_dir)

    spec = importlib.util.spec_from_file_location("peerpdf_garanti_dipnot", PEERPDF_SCRIPT_PATH)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"PeerPDF Garanti parser yuklenemedi: {PEERPDF_SCRIPT_PATH}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def extract_peerpdf_tables(pdf_path: Path) -> list[tuple[str, list[str], list[list[str]]]]:
    parser = load_peerpdf_garanti_module()
    tables = []

    for title in parser.CREDIT_TABLE_TITLES:
        lines = parser.extract_table_lines(pdf_path, title)
        rows = parser.parse_rows(lines)
        tables.append((title, parser.CREDIT_HEADERS, rows))

    tables.append(
        (
            parser.DEPOSIT_MATURITY_TITLE,
            parser.DEPOSIT_HEADERS,
            parser.extract_deposit_maturity_rows(pdf_path),
        )
    )
    tables.append(
        (
            parser.DEPOSIT_TOTAL_BY_MATURITY_TITLE,
            parser.DEPOSIT_TOTAL_HEADERS,
            parser.extract_deposit_total_by_maturity_rows(pdf_path),
        )
    )
    tables.extend(extract_deposit_interest_tables(pdf_path))
    tables.append(
        (
            parser.RESTRUCTURED_NPL_TITLE,
            parser.RESTRUCTURED_NPL_HEADERS,
            parser.extract_restructured_npl_rows(pdf_path),
        )
    )
    tables.append(
        (
            parser.NPL_MOVEMENT_TITLE,
            parser.NPL_MOVEMENT_HEADERS,
            parser.extract_npl_movement_rows(pdf_path),
        )
    )
    tables.append(
        (
            parser.LOAN_INTEREST_INCOME_TITLE,
            parser.LOAN_INTEREST_INCOME_HEADERS,
            parser.extract_loan_interest_income_rows(pdf_path),
        )
    )
    tables.append(
        (
            parser.SECURITIES_INTEREST_INCOME_TITLE,
            parser.SECURITIES_INTEREST_INCOME_HEADERS,
            parser.extract_securities_interest_income_rows(pdf_path),
        )
    )
    tables.append(
        (
            parser.EQUITY_ITEMS_TITLE,
            parser.EQUITY_ITEMS_HEADERS,
            parser.extract_equity_item_rows(pdf_path),
        )
    )
    tables.append(
        (
            parser.EXPECTED_LOSS_PROVISIONS_TITLE,
            parser.EXPECTED_LOSS_PROVISIONS_HEADERS,
            parser.extract_expected_loss_provision_rows(pdf_path),
        )
    )
    tables.append(
        (
            parser.STAGE3_PROVISIONS_TITLE,
            parser.STAGE3_PROVISIONS_HEADERS,
            parser.extract_stage3_provision_rows(pdf_path),
        )
    )
    tables.append(
        (
            parser.OTHER_OPERATING_INCOME_TITLE,
            parser.OTHER_OPERATING_INCOME_HEADERS,
            parser.extract_other_operating_income_rows(pdf_path),
        )
    )
    tables.append(
        (
            parser.CASH_AND_TCMB_TITLE,
            parser.CASH_AND_TCMB_HEADERS,
            parser.extract_cash_and_tcmb_rows(pdf_path),
        )
    )
    tables.append(
        (
            parser.USED_LOANS_INTEREST_EXPENSE_TITLE,
            parser.USED_LOANS_INTEREST_EXPENSE_HEADERS,
            parser.extract_used_loans_interest_expense_rows(pdf_path),
        )
    )
    tables.append(
        (
            parser.BANKS_INTEREST_INCOME_TITLE,
            parser.BANKS_INTEREST_INCOME_HEADERS,
            parser.extract_banks_interest_income_rows(pdf_path),
        )
    )
    tables.append(
        (
            parser.LIQUIDITY_COVERAGE_RATIO_TITLE,
            parser.LIQUIDITY_COVERAGE_RATIO_HEADERS,
            parser.extract_liquidity_coverage_ratio_rows(pdf_path),
        )
    )
    tables.append(extract_stage_provision_rows(pdf_path))
    return tables


def materialize_pdf_payload(pdf_path: Path) -> Path:
    with pdf_path.open("rb") as handle:
        header = handle.read(32)
    if header.startswith(b"%PDF-"):
        return pdf_path
    payload = pdf_path.read_bytes()
    start = payload.find(b"%PDF-")
    end = payload.rfind(b"%%EOF")
    if start < 0 or end < start:
        raise RuntimeError(f"Dosya PDF payload icermiyor: {pdf_path}")
    end += len(b"%%EOF")
    digest = hashlib.sha1(str(pdf_path.resolve()).encode("utf-8")).hexdigest()[:12]
    target_dir = Path(tempfile.gettempdir()) / "peer_analysis_pdf_payloads"
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / f"{digest}_{pdf_path.name}"
    raw_pdf = payload[start:end] + b"\n"
    if not target.exists() or target.stat().st_size != len(raw_pdf):
        target.write_bytes(raw_pdf)
    return target


def candidate_page_numbers(
    title: str,
    page_texts: list[str],
    parser,
    include_next_page: bool = True,
) -> list[int] | None:
    """Return 1-based pages containing a title, plus continuation pages."""
    matched = [
        index
        for index, page_text in enumerate(page_texts, start=1)
        if parser.title_in_page(title, page_text)
    ]
    if not matched:
        return None
    selected: set[int] = set(matched)
    if include_next_page:
        selected.update(
            page_number + 1
            for page_number in matched
            if page_number < len(page_texts)
        )
    return sorted(selected)


def _run_with_limited_pdf_pages(parser, page_numbers: list[int] | None, callback):
    if not page_numbers:
        return callback()
    original_open = parser.pdfplumber.open

    def limited_open(*args, **kwargs):
        kwargs.setdefault("pages", page_numbers)
        return original_open(*args, **kwargs)

    parser.pdfplumber.open = limited_open
    try:
        return callback()
    finally:
        parser.pdfplumber.open = original_open


def safe_extract_peerpdf_tables(
    pdf_path: Path,
    use_per_extractor_timeout: bool = True,
) -> list[tuple[str, list[str], list[list[str]]]]:
    source_pdf_path = Path(pdf_path)
    pdf_path = materialize_pdf_payload(source_pdf_path)
    parser = load_peerpdf_garanti_module()
    page_texts = load_or_build_pdf_page_texts(source_pdf_path, pdf_path)
    extractors: list[tuple[str, list[str], str, str | None]] = []
    for title in parser.CREDIT_TABLE_TITLES:
        extractors.append((title, parser.CREDIT_HEADERS, "credit", title))
    extractors.extend(
        [
            (parser.DEPOSIT_MATURITY_TITLE, parser.DEPOSIT_HEADERS, "parser_function", "extract_deposit_maturity_rows"),
            (parser.DEPOSIT_TOTAL_BY_MATURITY_TITLE, parser.DEPOSIT_TOTAL_HEADERS, "parser_function", "extract_deposit_total_by_maturity_rows"),
            (parser.RESTRUCTURED_NPL_TITLE, parser.RESTRUCTURED_NPL_HEADERS, "parser_function", "extract_restructured_npl_rows"),
            (parser.NPL_MOVEMENT_TITLE, parser.NPL_MOVEMENT_HEADERS, "parser_function", "extract_npl_movement_rows"),
            (parser.LOAN_INTEREST_INCOME_TITLE, parser.LOAN_INTEREST_INCOME_HEADERS, "parser_function", "extract_loan_interest_income_rows"),
            (parser.SECURITIES_INTEREST_INCOME_TITLE, parser.SECURITIES_INTEREST_INCOME_HEADERS, "parser_function", "extract_securities_interest_income_rows"),
            (parser.EQUITY_ITEMS_TITLE, parser.EQUITY_ITEMS_HEADERS, "parser_function", "extract_equity_item_rows"),
            (parser.EXPECTED_LOSS_PROVISIONS_TITLE, parser.EXPECTED_LOSS_PROVISIONS_HEADERS, "parser_function", "extract_expected_loss_provision_rows"),
            (parser.STAGE3_PROVISIONS_TITLE, parser.STAGE3_PROVISIONS_HEADERS, "parser_function", "extract_stage3_provision_rows"),
            (parser.OTHER_OPERATING_INCOME_TITLE, parser.OTHER_OPERATING_INCOME_HEADERS, "parser_function", "extract_other_operating_income_rows"),
            (parser.OTHER_OPERATING_EXPENSE_TITLE, parser.OTHER_OPERATING_EXPENSE_HEADERS, "parser_function", "extract_other_operating_expense_rows"),
            (parser.FX_RISK_TITLE, parser.FX_RISK_HEADERS, "parser_function", "extract_fx_risk_rows"),
            (parser.CASH_AND_TCMB_TITLE, parser.CASH_AND_TCMB_HEADERS, "parser_function", "extract_cash_and_tcmb_rows"),
            (parser.USED_LOANS_INTEREST_EXPENSE_TITLE, parser.USED_LOANS_INTEREST_EXPENSE_HEADERS, "parser_function", "extract_used_loans_interest_expense_rows"),
            (parser.BANKS_INTEREST_INCOME_TITLE, parser.BANKS_INTEREST_INCOME_HEADERS, "parser_function", "extract_banks_interest_income_rows"),
            (parser.LIQUIDITY_COVERAGE_RATIO_TITLE, parser.LIQUIDITY_COVERAGE_RATIO_HEADERS, "parser_function", "extract_liquidity_coverage_ratio_rows"),
        ]
    )

    period = extract_report_date(pdf_path) or "DONEM_YOK"
    tables: list[tuple[str, list[str], list[list[str]]]] = []

    def run_one_extractor(item):
        title, headers, extractor_kind, extractor_arg = item
        started = time.monotonic()
        page_numbers = candidate_page_numbers(title, page_texts, parser)
        page_description = (
            ",".join(str(number) for number in page_numbers)
            if page_numbers
            else "tam PDF yedegi"
        )
        print(
            f"PDF dipnot baslik deneniyor: {period} - {title} "
            f"(sayfa: {page_description})",
            flush=True,
        )
        if use_per_extractor_timeout:
            rows = run_extractor_with_timeout(
                pdf_path,
                extractor_kind,
                extractor_arg,
                page_numbers,
            )
        else:
            rows = extract_table_worker(
                str(pdf_path),
                extractor_kind,
                extractor_arg,
                page_numbers,
            )
        return title, headers, rows, time.monotonic() - started

    extractor_results: dict[str, tuple[list[str], list[list[str]], float]] = {}
    worker_count = min(CURRENT_PDF_EXTRACTOR_WORKERS, len(extractors))
    print(
        f"PDF dipnotlari {worker_count} paralel isciyle ayrisitiliyor.",
        flush=True,
    )
    with concurrent.futures.ThreadPoolExecutor(max_workers=worker_count) as executor:
        future_map = {
            executor.submit(run_one_extractor, item): item
            for item in extractors
        }
        for future in concurrent.futures.as_completed(future_map):
            title, headers, _extractor_kind, _extractor_arg = future_map[future]
            extractor_started = time.monotonic()
            try:
                result_title, result_headers, rows, elapsed = future.result()
                extractor_results[result_title] = (result_headers, rows, elapsed)
                print(
                    f"PDF dipnot baslik tamamlandi: {period} - {result_title} -> "
                    f"{len(rows)} satir, {elapsed:.1f} sn",
                    flush=True,
                )
            except concurrent.futures.TimeoutError:
                print(
                    f"WARN: {period} dipnot alinamadi: {title} -> "
                    f"Timeout: {EXTRACTOR_TIMEOUT_SECONDS} sn asildi",
                    flush=True,
                )
            except Exception as exc:
                print(
                    f"WARN: {period} dipnot alinamadi: {title} -> "
                    f"{type(exc).__name__}: {exc}",
                    flush=True,
                )

    # Paralel tamamlama sırası değişken olsa da Excel tablo sırası sabit kalsın.
    for title, _headers, _extractor_kind, _extractor_arg in extractors:
        result = extractor_results.get(title)
        if result is None:
            continue
        headers, rows, _elapsed = result
        tables.append((title, headers, rows))

    deposit_interest_pages = candidate_page_numbers(
        "Mevduata odenen faizin vade yapisina gore gosterimi",
        page_texts,
        parser,
    )
    stage_provision_pages = candidate_page_numbers(
        STAGE_PROVISION_TITLE,
        page_texts,
        parser,
    )

    try:
        print(f"PDF dipnot baslik deneniyor: {period} - Mevduata odenen faizin vade yapisi", flush=True)
        tables.extend(
            _run_with_limited_pdf_pages(
                parser,
                deposit_interest_pages,
                lambda: extract_deposit_interest_tables(pdf_path),
            )
        )
    except Exception as exc:
        print(f"WARN: {period} dipnot alinamadi: Mevduata odenen faizin vade yapisi -> {type(exc).__name__}: {exc}", flush=True)

    try:
        print(f"PDF dipnot baslik deneniyor: {period} - {STAGE_PROVISION_TITLE}", flush=True)
        tables.append(
            _run_with_limited_pdf_pages(
                parser,
                stage_provision_pages,
                lambda: extract_stage_provision_rows(pdf_path),
            )
        )
    except Exception as exc:
        print(f"WARN: {period} dipnot alinamadi: {STAGE_PROVISION_TITLE} -> {type(exc).__name__}: {exc}", flush=True)

    tables = apply_garanti_text_fallbacks(tables, page_texts, parser)

    # Three Garanti layouts expose most tables as positioned text instead of
    # pdfplumber grids. Reuse the common text parser only for sections the
    # Garanti parser could not produce; never overwrite a Garanti result.
    if len(tables) < 15:
        try:
            fallback_path = Path(__file__).resolve().parents[1] / "teb_bankasi" / "peerpdf_scripts" / "teb_dipnot.py"
            spec = importlib.util.spec_from_file_location("garanti_common_text_fallback", fallback_path)
            if spec is None or spec.loader is None:
                raise RuntimeError("Ortak metin parser yuklenemedi")
            fallback = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(fallback)
            existing = {_section_key(title) for title, _headers, _rows in tables}
            for title, headers, rows in fallback.extract_all_tables(pdf_path):
                key = _section_key(title)
                if key and key not in existing:
                    tables.append((title, headers, rows))
                    existing.add(key)
        except Exception as exc:
            print(f"WARN: {period} ortak metin yedegi calismadi -> {type(exc).__name__}: {exc}", flush=True)

    return tables


def _numbers_after_label(line: str, label_key: str) -> list[str]:
    if label_key not in normalize_key(line):
        return []
    return re.findall(r"(?:-|\d+(?:,\d{3})*)", line)


def _replace_semantic_table(tables, key: str, table) -> None:
    for index, (title, _headers, _rows) in enumerate(tables):
        if _section_key(title) == key:
            tables[index] = table
            return
    tables.append(table)


def apply_garanti_text_fallbacks(tables, page_texts: list[str], parser):
    """Recover Garanti Q2 positioned-text tables missed by grid extraction."""
    # Bank interest income: the current TP/YP figures are the first pair.
    bank_rows: list[list[str]] = []
    for text in page_texts:
        text_key = normalize_key(text)
        # pdfplumber'ın bu raporda bazı Türkçe harfleri U+FFFD olarak
        # döndürmesi nedeniyle başlığı sağlam kalan sözcüklerle tanıyoruz.
        if "bankalardan al" not in text_key or "faiz gelirlerine" not in text_key:
            continue
        active = False
        yurt_bank_index = 0
        for line in text.splitlines():
            key = normalize_key(line)
            if "bankalardan al" in key and "faiz gelirlerine" in key:
                active = True
                continue
            if active and "menkul de" in key and "faiz gelirlerine" in key:
                break
            if not active:
                continue
            label = None
            if "merkez bankas" in key:
                label = "T.C. Merkez Bankasından"
            elif "yurt" in key and "bankalardan" in key:
                label = "Yurt İçi Bankalardan" if yurt_bank_index == 0 else "Yurt Dışı Bankalardan"
                yurt_bank_index += 1
            elif "yurt" in key and "ubelerden" in key:
                label = "Yurt Dışı Merkez ve Şubelerden"
            elif key.startswith("toplam"):
                label = "Toplam"
            if label is not None:
                tokens = re.findall(r"(?:-|\d+(?:,\d{3})*)", line)
                if len(tokens) >= 2:
                    bank_rows.append([label, tokens[0], tokens[1]])
        if bank_rows:
            break
    if len(bank_rows) >= 4:
        _replace_semantic_table(
            tables,
            "banks_interest",
            (parser.BANKS_INTEREST_INCOME_TITLE, ["Kalem", "TP", "YP"], bank_rows),
        )

    equity_targets = (
        (("ekirdek sermaye toplam",), "Çekirdek Sermaye Toplamı", False),
        (("ndirimler", "ncesi katk", "sermaye"), "Katkı Sermaye Toplamı", False),
        (("toplam", "zkaynak", "ana sermaye", "katk", "sermaye toplam"), "Toplam Özkaynak (Ana sermaye ve katkı sermaye toplamı)", True),
        (("toplam risk a", "tutarlar"), "Toplam Risk Ağırlıklı Tutarlar", False),
    )
    equity_rows: list[list[str]] = []
    for fragments, label, take_last in equity_targets:
        value = None
        for text in page_texts:
            for line in text.splitlines():
                line_key = normalize_key(line)
                if all(fragment in line_key for fragment in fragments):
                    tokens = re.findall(r"\d+(?:,\d{3})+|\d+", line)
                    if tokens:
                        value = tokens[0]
                        if not take_last:
                            break
            if value is not None and not take_last:
                break
        if value is not None:
            equity_rows.append([label, "", "", value])
    if len(equity_rows) == len(equity_targets):
        _replace_semantic_table(
            tables,
            "equity",
            (parser.EQUITY_ITEMS_TITLE, ["Kalem", "", "", "Toplam"], equity_rows),
        )

    summary_text = next(
        (text for text in page_texts if "1. asama nakdi krediler" in normalize_key(text) and "3. asama donuk alacak" in normalize_key(text)),
        "",
    )
    if summary_text:
        compact = " ".join(summary_text.split())
        summary_specs = (
            (r"1\. Aşama Nakdi Krediler\s+((?:[\d,-]+\s+){5}[\d,-]+)", "1. Aşama Nakdi Krediler"),
            (r"2\. Aşama Nakdi Krediler\s+((?:[\d,-]+\s+){5}[\d,-]+)", "2. Aşama Nakdi Krediler"),
            (r"1\. ve 2\. Aşama Krediler\s+((?:[\d,-]+\s+){5}[\d,-]+)\s+Toplamı", "1. ve 2. Aşama Krediler Toplamı"),
            (r"1\. ve 2\. Aşama Karşılıkları\s+((?:[\d,-]+\s+){5}[\d,-]+)", "1. ve 2. Aşama Karşılıkları"),
            (r"3\. Aşama Donuk Alacaklar\s+((?:[\d,-]+\s+){5}[\d,-]+)\s+Toplamı", "3. Aşama Donuk Alacaklar Toplamı"),
            (r"3\. Aşama Donuk Alacak\s+((?:[\d,-]+\s+){5}[\d,-]+)\s+Karşılıkları", "3. Aşama Donuk Alacak Karşılıkları"),
        )
        summary_rows: list[list[str]] = []
        for pattern, label in summary_specs:
            match = re.search(pattern, compact, re.IGNORECASE)
            if not match:
                continue
            values = match.group(1).split()
            summary_rows.append([label, values[-2], values[-1]])
        if len(summary_rows) == len(summary_specs):
            tables.append(
                (
                    "Birinci ve ikinci grup krediler, diğer alacaklar ile yeniden yapılandırılan ya da yeni bir itfa planına bağlanan krediler ve bunlar için ayrılan karşılıklara ilişkin bilgiler",
                    ["Kalem", "TL", "YP"],
                    summary_rows,
                )
            )
    return tables


def run_extractor_with_timeout(
    pdf_path: Path,
    extractor_kind: str,
    extractor_arg: str | None,
    page_numbers: list[int] | None = None,
) -> list[list[str]]:
    context = multiprocessing.get_context("spawn")
    queue = context.Queue()
    process = context.Process(
        target=extract_table_worker_process,
        args=(queue, str(pdf_path), extractor_kind, extractor_arg, page_numbers),
    )
    process.start()
    process.join(EXTRACTOR_TIMEOUT_SECONDS)
    if process.is_alive():
        process.terminate()
        process.join(5)
        if process.is_alive():
            process.kill()
            process.join(5)
        raise concurrent.futures.TimeoutError()
    if queue.empty():
        raise RuntimeError(f"Extractor sonuc uretmedi, exitcode={process.exitcode}")
    status, payload = queue.get()
    if status == "ok":
        return payload
    raise RuntimeError(payload)


def extract_table_worker_process(
    queue,
    pdf_path_text: str,
    extractor_kind: str,
    extractor_arg: str | None,
    page_numbers: list[int] | None = None,
) -> None:
    try:
        queue.put(
            (
                "ok",
                extract_table_worker(
                    pdf_path_text,
                    extractor_kind,
                    extractor_arg,
                    page_numbers,
                ),
            )
        )
    except Exception as exc:
        queue.put(("err", f"{type(exc).__name__}: {exc}"))


def extract_table_worker(
    pdf_path_text: str,
    extractor_kind: str,
    extractor_arg: str | None,
    page_numbers: list[int] | None = None,
) -> list[list[str]]:
    pdf_path = Path(pdf_path_text)
    parser = load_peerpdf_garanti_module()

    def execute():
        if extractor_kind == "credit":
            if not extractor_arg:
                raise RuntimeError("Credit extractor title bos.")
            return parser.parse_rows(
                parser.extract_table_lines(pdf_path, extractor_arg)
            )
        if extractor_kind == "parser_function":
            if not extractor_arg:
                raise RuntimeError("Parser function adi bos.")
            extractor = getattr(parser, extractor_arg)
            return extractor(pdf_path)
        raise RuntimeError(f"Bilinmeyen extractor tipi: {extractor_kind}")

    return _run_with_limited_pdf_pages(parser, page_numbers, execute)


def parse_amount_text(value: str | None) -> int | None:
    if value is None:
        return None
    cleaned = " ".join(str(value).replace("\n", " ").split())
    if cleaned in {"", "-"}:
        return None
    digits = "".join(char for char in cleaned if char.isdigit())
    if not digits:
        return None
    return int(digits)


def format_amount(value: int | None) -> str:
    if value is None:
        return "-"
    return f"{value:,}"


def extract_deposit_interest_tables(pdf_path: Path) -> list[tuple[str, list[str], list[list[str]]]]:
    helper_path = PEERPDF_SCRIPT_PATH.parent / "mevduat_faiz.py"
    spec = importlib.util.spec_from_file_location("peerpdf_deposit_interest", helper_path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"PeerPDF mevduat faiz parser yuklenemedi: {helper_path}")
    helper = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(helper)

    extracted = helper.extract_tables(pdf_path)
    if not extracted:
        raise RuntimeError(f"Hedef tablo PDF icinde bulunamadi: {helper.TABLE_TITLE}")

    _, rows = extracted[0]
    summary_rows: list[list[str]] = []
    for row in rows:
        if not row:
            continue
        label = helper.clean_cell(row[0])
        vadesiz = helper.clean_cell(row[1] if len(row) > 1 else "")
        toplam = helper.clean_cell(row[-1] if len(row) > 1 else "")
        vadesiz_amount = parse_amount_text(vadesiz)
        toplam_amount = parse_amount_text(toplam)
        vadeli = (
            None
            if toplam_amount is None
            else toplam_amount - (vadesiz_amount or 0)
        )
        summary_rows.append([label, vadesiz, format_amount(vadeli), toplam])

    total_row = next(
        (row for row in rows if row and helper.clean_cell(row[0]) == "Genel Toplam"),
        rows[-1],
    )
    maturity_headers = [
        helper.clean_cell(header if header != "{report_date}" else "Vadesiz Mevduat")
        for header in helper.HEADERS[1:]
    ]
    transposed_rows = [
        [header, "", "", helper.clean_cell(value)]
        for header, value in zip(maturity_headers, total_row[1:])
    ]

    return [
        (helper.TABLE_TITLE, ["Kalem", "Vadesiz", "Vadeli", "Toplam"], summary_rows),
        (
            f"{helper.TABLE_TITLE} - toplam vade kirilimi",
            ["Vade", "", "", "Toplam"],
            transposed_rows,
        ),
    ]


def extract_stage_provision_rows(pdf_path: Path) -> tuple[str, list[str], list[list[str]]]:
    parser = load_peerpdf_garanti_module()
    helper_path = PEERPDF_SCRIPT_PATH.parent / "birinci_ikinci_asama_krediler.py"
    spec = importlib.util.spec_from_file_location("peerpdf_stage_provisions", helper_path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"PeerPDF kredi asama parser yuklenemedi: {helper_path}")
    helper = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(helper)

    target_title = normalize_key(helper.TABLE_TITLE)
    target_labels = {normalize_key(label) for label in STAGE_PROVISION_TARGET_LABELS}

    with parser.pdfplumber.open(str(pdf_path)) as pdf:
        start_page_index = None
        for page_index, page in enumerate(pdf.pages):
            text = page.extract_text(x_tolerance=1, y_tolerance=3) or ""
            if target_title not in normalize_key(text) and not parser.title_in_page(helper.TABLE_TITLE, text):
                continue
            start_page_index = page_index
            break

        if start_page_index is None:
            raise RuntimeError(f"Hedef baslik PDF icinde bulunamadi: {helper.TABLE_TITLE}")

        for page in pdf.pages[start_page_index : min(len(pdf.pages), start_page_index + 4)]:
            for table in page.extract_tables():
                if not table or len(table) < 3:
                    continue
                rows: list[list[str]] = []
                for raw_row in table:
                    if not raw_row or not raw_row[0]:
                        continue
                    label = helper.clean_cell(raw_row[0])
                    if normalize_key(label) not in target_labels:
                        continue
                    rows.append(
                        [
                            label,
                            helper.clean_cell(raw_row[1]),
                            helper.clean_cell(raw_row[2]),
                            helper.clean_cell(raw_row[3]),
                            helper.clean_cell(raw_row[4]),
                        ]
                    )
                if rows:
                    return (
                        helper.TABLE_TITLE,
                        [
                            "Kalem",
                            "Cari Donem - Standart Nitelikli Krediler",
                            "Cari Donem - Yakin Izlemedeki Krediler",
                            "Onceki Donem - Standart Nitelikli Krediler",
                            "Onceki Donem - Yakin Izlemedeki Krediler",
                        ],
                        rows,
                    )

        for page in pdf.pages[start_page_index : min(len(pdf.pages), start_page_index + 4)]:
            text = page.extract_text(x_tolerance=1, y_tolerance=3) or ""
            found: list[list[str]] = []
            for raw_line in text.splitlines():
                line = " ".join(raw_line.split())
                line_key = normalize_key(line)
                label = next((x for x in STAGE_PROVISION_TARGET_LABELS if normalize_key(x) in line_key), None)
                if label is None:
                    continue
                tokens = re.findall(r"(?:-|\d+(?:,\d{3})*)", line)
                if normalize_key(label).startswith("12 aylik") and tokens and tokens[0] == "12":
                    tokens = tokens[1:]
                if normalize_key(label).startswith("12 aylik") and len(tokens) == 3:
                    tokens = [tokens[0], "-", tokens[1], tokens[2]]
                elif normalize_key(label).startswith("kredi riskinde") and len(tokens) == 3:
                    tokens = ["-", tokens[0], tokens[1], tokens[2]]
                if len(tokens) >= 4:
                    found.append([label, *tokens[:4]])
            if len(found) >= 2:
                return (
                    helper.TABLE_TITLE,
                    ["Kalem", "Cari Donem - Standart Nitelikli Krediler", "Cari Donem - Yakin Izlemedeki Krediler", "Onceki Donem - Standart Nitelikli Krediler", "Onceki Donem - Yakin Izlemedeki Krediler"],
                    found[:2],
                )

    raise RuntimeError(f"Hedef iki karsilik satiri PDF icinde bulunamadi: {helper.TABLE_TITLE}")


def normalize_text(text: str) -> str:
    text = text.replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def parse_pdf_notes(text: str) -> list[PdfNote]:
    normalized = normalize_text(text)
    pattern = re.compile(
        r"(?is)(?:dipnot|not)\s+(\d+)\s*[-.:)]?\s*([^\n]{0,180})\n(.*?)(?=\n\s*(?:dipnot|not)\s+\d+\s*[-.:)]?|\Z)"
    )
    notes: list[PdfNote] = []
    for match in pattern.finditer(normalized):
        note_no, title, body = match.groups()
        cleaned_title = " ".join(title.split())
        cleaned_body = normalize_text(body)
        if cleaned_body:
            notes.append(PdfNote(note_no=note_no, title=cleaned_title, text=cleaned_body))

    if notes:
        return notes

    excerpt = normalized[:MAX_CELL_TEXT_LENGTH]
    return [PdfNote(note_no="", title="PDF metni", text=excerpt)]


def trim_cell_text(value: str) -> str:
    return value[:MAX_CELL_TEXT_LENGTH]


def delete_existing_pdf_block(sheet) -> None:
    for row in range(1, sheet.max_row + 1):
        if (
            sheet.cell(row=row, column=1).value == BLOCK_TITLE
            or sheet.cell(row=row, column=26).value == BLOCK_TITLE
            or sheet.cell(row=row, column=MARKER_COLUMN).value == BLOCK_TITLE
        ):
            sheet.delete_rows(row, sheet.max_row - row + 1)
            return


def append_peerpdf_tables_to_sheet(
    workbook_path: Path,
    sheet_name: str,
    pdf_path: Path,
    tables: list[tuple[str, list[str], list[list[str]]]],
) -> int:
    workbook = load_workbook(workbook_path)
    if sheet_name not in workbook.sheetnames:
        raise RuntimeError(f"Workbook icinde sheet bulunamadi: {sheet_name}")

    sheet = workbook[sheet_name]
    delete_existing_pdf_block(sheet)
    report_date = extract_report_date(pdf_path)
    value_start_column = period_start_column(sheet, report_date)

    current_row = sheet.max_row + 3
    sheet.cell(row=current_row, column=MARKER_COLUMN, value=BLOCK_TITLE)
    sheet.column_dimensions[get_column_letter(MARKER_COLUMN)].hidden = True
    total_rows = 0
    for table_title, headers, rows in tables:
        if normalize_key(table_title) == normalize_key(STAGE_PROVISION_TITLE):
            current_row, written_rows = append_stage_provision_table(
                sheet=sheet,
                current_row=current_row,
                table_title=table_title,
                headers=headers,
                rows=rows,
                current_start_column=value_start_column,
            )
            total_rows += written_rows
            continue

        sheet.cell(row=current_row, column=1, value=table_title)
        sheet.cell(row=current_row, column=1).font = Font(bold=True, size=12)
        sheet.cell(row=current_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        current_row += 1

        sheet.cell(row=current_row, column=1, value=headers[0] if headers else "Kalem")
        sheet.cell(row=current_row, column=1).font = Font(bold=True)
        sheet.cell(row=current_row, column=1).fill = PatternFill("solid", fgColor=SOURCE_FILL)
        sheet.cell(row=current_row, column=1).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for offset, header in enumerate(headers[1:], start=0):
            cell = sheet.cell(row=current_row, column=value_start_column + offset, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor=SOURCE_FILL)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        current_row += 1

        for row in rows:
            if row:
                label_cell = sheet.cell(row=current_row, column=1, value=row[0])
                label_cell.alignment = Alignment(vertical="top", wrap_text=True)
            for offset, value in enumerate(row[1:], start=0):
                cell = sheet.cell(row=current_row, column=value_start_column + offset, value=value)
                cell.alignment = Alignment(vertical="top")
                if isinstance(value, int):
                    cell.number_format = "#,##0"
            current_row += 1
            total_rows += 1
        current_row += 1

    for column_index in [1, *range(value_start_column, min(sheet.max_column, value_start_column + 4) + 1)]:
        column_letter = get_column_letter(column_index)
        if column_index == 1:
            sheet.column_dimensions[column_letter].width = max(sheet.column_dimensions[column_letter].width or 0, 48)
        else:
            sheet.column_dimensions[column_letter].width = max(sheet.column_dimensions[column_letter].width or 0, 18)

    workbook.save(workbook_path)
    return total_rows


def _repair_mojibake_label(value: object) -> str:
    text = str(value or "")
    if any(marker in text for marker in ("Ã", "Ä", "Å")):
        try:
            text = text.encode("cp1252").decode("utf-8")
        except (UnicodeEncodeError, UnicodeDecodeError):
            pass
    return text


def _stable_row_label(value: object) -> str:
    text = _repair_mojibake_label(value)
    key = normalize_key(text)
    key = re.sub(r"\s*\(\s*\*+\s*\)\s*$", "", key)
    key = re.sub(r"\b(?:dipnot|aciklama)\s*\d+\b", "", key)
    key = re.sub(r"\b\d+\b$", "", key)
    if key in {"birikimli mevduat", "birikiml i", "birikimli"}:
        return "birikimli mevduat"
    if key.startswith("kayittan dusulen") or (
        "kayittan dusulen" in key and "aktiften silinen" in key
    ):
        return "kayittan dusulen"
    return " ".join(key.split())


def _section_key(title: str) -> str | None:
    key = normalize_key(title)
    if "tuketici kredileri" in key: return "consumer"
    if "taksitli ticari" in key: return "installment"
    if "mevduat" in key and "vade yapisi" in key and "oden" not in key: return "deposit_maturity_total" if "toplam bilgi" in key else "deposit_maturity"
    if "mevduat" in key and ("odenen faiz" in key or "odenen faizin" in key): return "deposit_interest_total" if "toplam vade" in key else "deposit_interest"
    if "toplam donuk alacak hareket" in key: return "npl_movement"
    if "donuk alacak" in key and ("yapilandir" in key or "faiz tahakkuk" in key): return "restructured"
    if "bankalardan alinan faiz" in key: return "banks_interest"
    if "menkul deger" in key and "faiz" in key: return "securities_interest"
    if "ozkaynak" in key: return "equity"
    if ("beklenen zarar" in key and "diger karsilik" in key) or "beklenen kredi zarari" in key or "deger dusus karsilik" in key: return "expected_loss"
    if ("standart nitelikli" in key and "yakin izle" in key) or "birinci ve ikinci asama" in key: return "stage12"
    if "likidite karsilama" in key: return "liquidity"
    if "ucuncu asama" in key or "3 asama" in key or "3. asama" in key or "ozel karsilik" in key: return "stage3"
    if "birinci ve ikinci grup krediler" in key and "ayrilan karsilik" in key: return "credit_summary"
    if "diger faaliyet gelir" in key: return "other_income"
    if "diger faaliyet gider" in key: return "other_expense"
    if "kur riskine iliskin aciklama" in key: return "fx_risk"
    if "nakit deger" in key and ("tcmb" in key or "merkez bank" in key): return "cash"
    if "kullanilan kredi" in key and "faiz" in key: return "used_loans"
    if "kredilerden alinan faiz" in key: return "loan_interest"
    return None


def _ordered_tables(tables_by_title: dict[str, dict[str, object]]) -> list[str]:
    by_key: dict[str, str] = {}
    extras: list[str] = []
    for title in tables_by_title:
        key = _section_key(title)
        if key and key not in by_key: by_key[key] = title
        else: extras.append(title)
    ordered: list[str] = []
    for _rank, _subrank, key, title, headers in TARGET_SECTIONS:
        actual = by_key.get(key)
        if actual is None:
            actual = title
            tables_by_title[actual] = {"headers": headers, "master_rows": []}
        ordered.append(actual)
    return ordered + extras


def _excel_number(value: object, percentage_like: bool = False) -> object:
    if isinstance(value, (int, float)) or value is None: return value
    text = str(value).strip().replace(" ", "")
    if text in {"", "-", "—", "–"}: return None
    negative = text.startswith("(") and text.endswith(")")
    if negative: text = text[1:-1]
    if not re.fullmatch(r"[+-]?\d[\d.,]*", text): return value
    if percentage_like:
        normalized = text.replace(",", ".")
        if normalized.count(".") > 1: normalized = normalized.replace(".", "")
        number: object = float(normalized)
    else:
        number = int(re.sub(r"[.,]", "", text))
    return -number if negative else number


def _row_label_score(left: str, right: str) -> int:
    if left == right and left:
        return 8
    if not left or not right:
        return -5
    critical = {"tp", "yp", "cari", "onceki", "standart", "toplam"}
    if ({*left.split()} & critical) != ({*right.split()} & critical):
        return -5
    ratio = SequenceMatcher(None, left, right).ratio()
    if ratio >= 0.92:
        return 5
    if ratio >= 0.78:
        return 2
    return -5


def _merge_period_rows(table_data: dict[str, object], period: str, rows: list[list[object]], stage_table: bool) -> None:
    master: list[dict[str, object]] = table_data["master_rows"]
    for existing in master:
        existing["label"] = _repair_mojibake_label(existing.get("label", ""))
        stable = _stable_row_label(existing.get("label", ""))
        if stable == "birikimli mevduat":
            existing["label"] = "Birikimli Mevduat"
        elif stable == "kayittan dusulen":
            existing["label"] = "Kayıttan Düşülen"
    incoming = [row for row in rows if row]
    for row in incoming:
        row[0] = _repair_mojibake_label(row[0])
        stable = _stable_row_label(row[0])
        if stable == "birikimli mevduat":
            row[0] = "Birikimli Mevduat"
        elif stable == "kayittan dusulen":
            row[0] = "Kayıttan Düşülen"
    deduped: list[list[object]] = []
    special_rows: dict[str, list[object]] = {}
    for row in incoming:
        stable = _stable_row_label(row[0])
        if stable not in {"birikimli mevduat", "kayittan dusulen"}:
            deduped.append(row)
            continue
        existing = special_rows.get(stable)
        if existing is None:
            special_rows[stable] = row
            deduped.append(row)
            continue
        for index in range(1, max(len(existing), len(row))):
            if index >= len(existing):
                existing.append(row[index])
            elif existing[index] in (None, "", "-") and index < len(row) and row[index] not in (None, "", "-"):
                existing[index] = row[index]
    incoming = deduped
    exact_other_count = sum(1 for row in incoming if _stable_row_label(row[0]) == "diger")
    table_data["max_exact_other_per_period"] = max(
        int(table_data.get("max_exact_other_per_period", 0)), exact_other_count,
    )
    if exact_other_count == 1:
        for row in incoming:
            if _stable_row_label(row[0]) == "diger":
                row[0] = "Diğer"
    if any("(cid:" in str(row[0]).casefold() for row in incoming):
        print(f"WARN: {period} tablosu bozuk PDF font kodlamasi (cid) nedeniyle atlandi; satir kaymasi engellendi.", flush=True)
        return
    master_keys = [_stable_row_label(row["label"]) for row in master]
    incoming_keys = [_stable_row_label(row[0]) for row in incoming]
    m, n = len(master_keys), len(incoming_keys)
    gap = -2
    dp = [[0] * (n + 1) for _ in range(m + 1)]
    trace = [[""] * (n + 1) for _ in range(m + 1)]
    for i in range(1, m + 1):
        dp[i][0], trace[i][0] = i * gap, "up"
    for j in range(1, n + 1):
        dp[0][j], trace[0][j] = j * gap, "left"
    for i in range(1, m + 1):
        for j in range(1, n + 1):
            choices = (
                (dp[i - 1][j - 1] + _row_label_score(master_keys[i - 1], incoming_keys[j - 1]), "diag"),
                (dp[i - 1][j] + gap, "up"),
                (dp[i][j - 1] + gap, "left"),
            )
            dp[i][j], trace[i][j] = max(choices, key=lambda item: item[0])
    operations: list[tuple[str, int | None, int | None]] = []
    i, j = m, n
    while i or j:
        op = trace[i][j]
        if op == "diag":
            operations.append((op, i - 1, j - 1)); i -= 1; j -= 1
        elif op == "up":
            operations.append((op, i - 1, None)); i -= 1
        else:
            operations.append(("left", None, j - 1)); j -= 1
    operations.reverse()
    rebuilt: list[dict[str, object]] = []
    for op, master_index, incoming_index in operations:
        if op == "up":
            rebuilt.append(master[master_index])
            continue
        row = incoming[incoming_index]
        if op == "diag" and _row_label_score(master_keys[master_index], incoming_keys[incoming_index]) > 0:
            target = master[master_index]
        elif op == "diag":
            rebuilt.append(master[master_index])
            target = {"label": str(row[0]), "values": {}}
        else:
            target = {"label": str(row[0]), "values": {}}
        values = row[1:3] if stage_table else row[1:]
        for offset, value in enumerate(values):
            target["values"][(period, offset)] = value
        rebuilt.append(target)
    if int(table_data.get("max_exact_other_per_period", 0)) <= 1:
        consolidated: list[dict[str, object]] = []
        other_target: dict[str, object] | None = None
        for row in rebuilt:
            if _stable_row_label(row["label"]) != "diger":
                consolidated.append(row)
                continue
            if other_target is None:
                row["label"] = "Diğer"
                other_target = row
                consolidated.append(row)
            else:
                other_target["values"].update(row["values"])
        rebuilt = consolidated
    table_data["master_rows"] = rebuilt


def append_peerpdf_history_to_sheet(
    workbook_path: Path,
    sheet_name: str,
    pdf_files: list[Path],
    cache_dir: Path | None = None,
    workers: int = DEFAULT_WORKERS,
) -> int:
    workbook = load_workbook(workbook_path)
    if sheet_name not in workbook.sheetnames:
        raise RuntimeError(f"Workbook icinde sheet bulunamadi: {sheet_name}")

    sheet = workbook[sheet_name]
    delete_existing_pdf_block(sheet)
    if sheet.max_row >= PDF_BLOCK_START_ROW:
        sheet.delete_rows(PDF_BLOCK_START_ROW, sheet.max_row - PDF_BLOCK_START_ROW + 1)

    table_order: list[str] = []
    tables_by_title: dict[str, dict[str, object]] = {}
    periods_written: set[str] = set()

    eligible_pdf_files: list[Path] = []
    for pdf_path in pdf_files:
        period = extract_report_date(pdf_path)
        if not period:
            print(f"WARN: PDF rapor donemi dosya adindan okunamadi, atlandi: {pdf_path}", flush=True)
            continue
        try:
            period_start_column(sheet, period)
        except Exception as exc:
            print(f"WARN: {period} icin workbook kolonlari bulunamadi, PDF atlandi: {pdf_path} -> {exc}", flush=True)
            continue
        eligible_pdf_files.append(pdf_path)

    extracted_by_pdf = load_history_pdf_tables(eligible_pdf_files, cache_dir, workers)

    for pdf_path in eligible_pdf_files:
        period = extract_report_date(pdf_path)
        tables = extracted_by_pdf.get(str(pdf_path.resolve()), [])
        if not tables:
            print(f"WARN: {period} icin hic PDF dipnot tablosu alinamadi: {pdf_path}", flush=True)
            continue
        periods_written.add(period)

        for table_title, headers, rows in tables:
            semantic_key = _section_key(table_title)
            existing_title = next(
                (known for known in tables_by_title if semantic_key and _section_key(known) == semantic_key),
                None,
            )
            if existing_title is not None:
                table_title = existing_title
            if table_title not in tables_by_title:
                tables_by_title[table_title] = {
                    "headers": headers,
                    "master_rows": [],
                }
                table_order.append(table_title)

            table_data = tables_by_title[table_title]
            _merge_period_rows(
                table_data, period, rows,
                normalize_key(table_title) == normalize_key(STAGE_PROVISION_TITLE),
            )

    if not periods_written:
        raise RuntimeError("History icinde islenebilir PDF dipnot verisi bulunamadi.")

    table_order = _ordered_tables(tables_by_title)

    current_row = PDF_BLOCK_START_ROW
    sheet.cell(row=current_row, column=MARKER_COLUMN, value=BLOCK_TITLE)
    sheet.column_dimensions[get_column_letter(MARKER_COLUMN)].hidden = True
    total_rows = 0

    for table_title in table_order:
        table_data = tables_by_title[table_title]
        headers = table_data["headers"]
        master_rows = table_data["master_rows"]

        sheet.cell(row=current_row, column=1, value=table_title)
        sheet.cell(row=current_row, column=1).font = Font(bold=True, size=12)
        sheet.cell(row=current_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        current_row += 1

        sheet.cell(row=current_row, column=1, value=headers[0] if headers else "Kalem")
        sheet.cell(row=current_row, column=1).font = Font(bold=True)
        sheet.cell(row=current_row, column=1).fill = PatternFill("solid", fgColor=SOURCE_FILL)
        sheet.cell(row=current_row, column=1).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for period in sorted(periods_written, key=period_sort_key):
            value_start_column = period_start_column(sheet, period)
            if normalize_key(table_title) == normalize_key(STAGE_PROVISION_TITLE):
                period_headers = ["Standart Nitelikli Krediler", "Yakin Izlemedeki Krediler"]
            elif _section_key(table_title) == "fx_risk":
                period_headers = ["Toplam"]
            else:
                period_headers = headers[1:]
            for offset, header in enumerate(period_headers):
                cell = sheet.cell(row=current_row, column=value_start_column + offset, value=header)
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor=SOURCE_FILL)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        current_row += 1

        for row_data in master_rows:
            label_cell = sheet.cell(row=current_row, column=1, value=row_data["label"])
            label_cell.alignment = Alignment(vertical="top", wrap_text=True)
            values = row_data["values"]
            for (period, offset), value in values.items():
                value_start_column = period_start_column(sheet, period)
                section_key = _section_key(table_title)
                numeric_value = _excel_number(value, section_key in {"liquidity", "fx_risk"})
                target_offset = 0 if section_key == "fx_risk" else offset
                cell = sheet.cell(row=current_row, column=value_start_column + target_offset, value=numeric_value)
                cell.alignment = Alignment(vertical="top")
                if isinstance(numeric_value, (int, float)):
                    if section_key == "fx_risk":
                        cell.number_format = "0.0000"
                    elif section_key == "liquidity":
                        cell.number_format = "0.00"
                    else:
                        cell.number_format = "#,##0;[Red](#,##0);-"
            current_row += 1
            total_rows += 1
        current_row += 1

    sheet.column_dimensions[get_column_letter(1)].width = max(sheet.column_dimensions[get_column_letter(1)].width or 0, 48)
    workbook.save(workbook_path)
    return total_rows


def load_history_pdf_tables(
    pdf_files: list[Path],
    cache_dir: Path | None,
    workers: int,
) -> dict[str, list[tuple[str, list[str], list[list[str]]]]]:
    if not pdf_files:
        return {}

    worker_count = max(1, min(workers, len(pdf_files)))
    if worker_count == 1:
        results: dict[str, list[tuple[str, list[str], list[list[str]]]]] = {}
        for index, pdf_path in enumerate(pdf_files, start=1):
            period = extract_report_date(pdf_path)
            print(f"PDF dipnot okunuyor ({index}/{len(pdf_files)}): {period} - {pdf_path.name}", flush=True)
            tables = load_or_extract_pdf_tables(pdf_path, cache_dir, use_per_extractor_timeout=False)
            print(f"PDF dipnot OK ({index}/{len(pdf_files)}): {period} - {len(tables)} tablo", flush=True)
            results[str(pdf_path.resolve())] = tables
        return results

    print(f"PDF dipnotlari paralel okunuyor: {len(pdf_files)} PDF, {worker_count} worker", flush=True)
    results = {}
    cache_dir_text = str(cache_dir) if cache_dir else ""
    with concurrent.futures.ProcessPoolExecutor(max_workers=worker_count) as executor:
        futures = {
            executor.submit(extract_pdf_tables_history_worker, str(pdf_path), cache_dir_text): pdf_path
            for pdf_path in pdf_files
        }
        completed = 0
        for future in concurrent.futures.as_completed(futures):
            pdf_path = futures[future]
            completed += 1
            period = extract_report_date(pdf_path)
            try:
                key, tables = future.result()
                results[key] = tables
                print(f"PDF dipnot OK ({completed}/{len(pdf_files)}): {period} - {len(tables)} tablo", flush=True)
            except Exception as exc:
                print(f"WARN: {period} PDF dipnotlari alinamadi: {pdf_path.name} -> {type(exc).__name__}: {exc}", flush=True)
                results[str(pdf_path.resolve())] = []
    return results


def extract_pdf_tables_history_worker(
    pdf_path_text: str,
    cache_dir_text: str,
) -> tuple[str, list[tuple[str, list[str], list[list[str]]]]]:
    pdf_path = Path(pdf_path_text)
    cache_dir = Path(cache_dir_text) if cache_dir_text else None
    print(f"PDF worker basladi: {extract_report_date(pdf_path)} - {pdf_path.name}", flush=True)
    tables = load_or_extract_pdf_tables(
        pdf_path,
        cache_dir,
        use_per_extractor_timeout=False,
    )
    return str(pdf_path.resolve()), tables


def cache_path_for_pdf(pdf_path: Path, cache_dir: Path) -> Path:
    period = extract_report_date(pdf_path) or "donem_yok"
    safe_stem = re.sub(r"[^A-Za-z0-9_.-]+", "_", pdf_path.stem).strip("_") or "pdf"
    return cache_dir / f"{period.replace('.', '')}_{safe_stem}.json"


def load_or_extract_pdf_tables(
    pdf_path: Path,
    cache_dir: Path | None,
    use_per_extractor_timeout: bool = True,
) -> list[tuple[str, list[str], list[list[str]]]]:
    if cache_dir is None:
        return sanitize_peerpdf_tables(
            safe_extract_peerpdf_tables(pdf_path, use_per_extractor_timeout=use_per_extractor_timeout)
        )

    cache_dir.mkdir(parents=True, exist_ok=True)
    cache_path = cache_path_for_pdf(pdf_path, cache_dir)
    pdf_stat = pdf_path.stat()
    if cache_path.exists():
        try:
            payload = json.loads(cache_path.read_text(encoding="utf-8"))
            if (
                payload.get("cache_version") == CACHE_VERSION
                and payload.get("source_size") == pdf_stat.st_size
                and payload.get("source_mtime") == pdf_stat.st_mtime
            ):
                tables = payload.get("tables") or []
                print(f"PDF cache kullanildi: {cache_path.name}", flush=True)
                return sanitize_peerpdf_tables([
                    (str(item["title"]), list(item["headers"]), [list(row) for row in item["rows"]])
                    for item in tables
                ])
        except Exception as exc:
            print(f"WARN: PDF cache okunamadi, yeniden ayrisitilacak: {cache_path} -> {type(exc).__name__}: {exc}", flush=True)

    tables = safe_extract_peerpdf_tables(pdf_path, use_per_extractor_timeout=use_per_extractor_timeout)
    payload = {
        "cache_version": CACHE_VERSION,
        "source_path": str(pdf_path),
        "source_size": pdf_stat.st_size,
        "source_mtime": pdf_stat.st_mtime,
        "tables": [
            {"title": title, "headers": headers, "rows": rows}
            for title, headers, rows in tables
        ],
    }
    cache_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"PDF cache yazildi: {cache_path.name}", flush=True)
    return sanitize_peerpdf_tables(tables)


def sanitize_peerpdf_tables(
    tables: list[tuple[str, list[str], list[list[str]]]]
) -> list[tuple[str, list[str], list[list[str]]]]:
    parser = load_peerpdf_garanti_module()
    credit_titles = {normalize_key(title) for title in parser.CREDIT_TABLE_TITLES}
    sanitized: list[tuple[str, list[str], list[list[str]]]] = []
    for title, headers, rows in tables:
        clean_rows = rows
        if normalize_key(title) in credit_titles:
            clean_rows = []
            for row in rows:
                if row:
                    row = [canonical_credit_label(str(row[0])), *row[1:]]
                clean_rows.append(row)
                if row and normalize_key(str(row[0])).startswith("toplam"):
                    break
        sanitized.append((title, headers, clean_rows))
    return sanitized


def canonical_credit_label(label: str) -> str:
    cleaned = " ".join(str(label).replace("\n", " ").split())
    replacements = {
        "T aşıt": "Taşıt",
        "T aksitli": "Taksitli",
        "T aksitsiz": "Taksitsiz",
        "Kre dileri": "Kredileri",
        "Kre disi": "Kredisi",
        "Kredili Mevduat Hesabı-TP(Personel)": "Kredili Mevduat Hesabı-TP (Personel)",
    }
    for old, new in replacements.items():
        cleaned = cleaned.replace(old, new)
    return cleaned


def append_stage_provision_table(
    sheet,
    current_row: int,
    table_title: str,
    headers: list[str],
    rows: list[list[str]],
    current_start_column: int,
) -> tuple[int, int]:
    sheet.cell(row=current_row, column=1, value=table_title)
    sheet.cell(row=current_row, column=1).font = Font(bold=True, size=12)
    sheet.cell(row=current_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    current_row += 1

    sheet.cell(row=current_row, column=1, value=headers[0] if headers else "Kalem")
    sheet.cell(row=current_row, column=1).font = Font(bold=True)
    sheet.cell(row=current_row, column=1).fill = PatternFill("solid", fgColor=SOURCE_FILL)
    sheet.cell(row=current_row, column=1).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    header_positions = [
        (current_start_column, "Standart Nitelikli Krediler"),
        (current_start_column + 1, "Yakin Izlemedeki Krediler"),
    ]
    for column_index, header in header_positions:
        cell = sheet.cell(row=current_row, column=column_index, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=SOURCE_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    current_row += 1

    written_rows = 0
    for row in rows:
        if not row:
            continue
        label_cell = sheet.cell(row=current_row, column=1, value=row[0])
        label_cell.alignment = Alignment(vertical="top", wrap_text=True)

        placements = [
            (current_start_column, row[1] if len(row) > 1 else ""),
            (current_start_column + 1, row[2] if len(row) > 2 else ""),
        ]
        for column_index, value in placements:
            cell = sheet.cell(row=current_row, column=column_index, value=value)
            cell.alignment = Alignment(vertical="top")
        current_row += 1
        written_rows += 1

    return current_row + 1, written_rows


def append_notes_to_sheet(workbook_path: Path, sheet_name: str, pdf_path: Path, notes: list[PdfNote]) -> int:
    workbook = load_workbook(workbook_path)
    if sheet_name not in workbook.sheetnames:
        raise RuntimeError(f"Workbook icinde sheet bulunamadi: {sheet_name}")

    sheet = workbook[sheet_name]
    delete_existing_pdf_block(sheet)

    start_row = sheet.max_row + 3
    sheet.cell(row=start_row, column=1, value=BLOCK_TITLE)
    sheet.cell(row=start_row, column=2, value=f"Kaynak: {pdf_path.name}")
    sheet.cell(row=start_row, column=3, value=f"Guncelleme: {datetime.now():%Y-%m-%d %H:%M:%S}")

    for cell in sheet[start_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    source_row = start_row + 1
    headers = ["Kaynak", "Dipnot", "Baslik", "Metin"]
    for column_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=source_row, column=column_index, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=SOURCE_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    current_row = source_row + 1
    for note in notes:
        sheet.cell(row=current_row, column=1, value=pdf_path.name)
        sheet.cell(row=current_row, column=2, value=note.note_no)
        sheet.cell(row=current_row, column=3, value=trim_cell_text(note.title))
        sheet.cell(row=current_row, column=4, value=trim_cell_text(note.text))
        for column_index in range(1, 5):
            sheet.cell(row=current_row, column=column_index).alignment = Alignment(vertical="top", wrap_text=True)
        current_row += 1

    sheet.column_dimensions[get_column_letter(4)].width = 90
    workbook.save(workbook_path)
    return len(notes)


def append_garanti_pdf_notes(
    workbook_path: Path,
    pdf_dir: Path,
    pdf_path: Path | None = None,
    sheet_name: str = COMPANY_NAME,
    cache_dir: Path | None = None,
    workers: int = DEFAULT_WORKERS,
) -> tuple[Path, int]:
    if pdf_path:
        source_pdf = pdf_path
        tables = safe_extract_peerpdf_tables(source_pdf)
        rows_written = append_peerpdf_tables_to_sheet(workbook_path, sheet_name, source_pdf, tables)
        return source_pdf, rows_written

    pdf_files = list_pdf_files(pdf_dir)
    if not pdf_files:
        raise FileNotFoundError(f"Garanti PDF dosyasi bulunamadi: {pdf_dir}")
    rows_written = append_peerpdf_history_to_sheet(workbook_path, sheet_name, pdf_files, cache_dir=cache_dir, workers=workers)
    return pdf_files[-1], rows_written


def main() -> None:
    parser = argparse.ArgumentParser(description="Garanti PDF dipnotlarini butunlesik Excel'deki Garanti sheetinin altina yazar.")
    parser.add_argument("--workbook", default=DEFAULT_WORKBOOK_PATH, help="Guncellenecek butunlesik Excel dosyasi.")
    parser.add_argument("--pdf-dir", default=DEFAULT_PDF_DIR, help="Garanti PDF dosyalarinin bulundugu klasor.")
    parser.add_argument("--pdf", help="Belirli bir PDF dosyasi kullan.")
    parser.add_argument("--cache-dir", default=DEFAULT_CACHE_DIR, help="PDF dipnot cache klasoru.")
    parser.add_argument("--workers", type=int, default=DEFAULT_WORKERS, help="History PDF ayrıştırma icin paralel worker sayisi.")
    parser.add_argument("--sheet", default=COMPANY_NAME, help="Guncellenecek sheet adi.")
    args = parser.parse_args()

    workbook_path = Path(args.workbook)
    pdf_dir = Path(args.pdf_dir)
    pdf_path = Path(args.pdf) if args.pdf else None
    cache_dir = Path(args.cache_dir) if args.cache_dir else None
    source_pdf, rows_written = append_garanti_pdf_notes(workbook_path, pdf_dir, pdf_path, args.sheet, cache_dir=cache_dir, workers=args.workers)
    print(f"PDF dipnotlari guncellendi: {source_pdf} -> {workbook_path} ({rows_written} satir)")


if __name__ == "__main__":
    main()
