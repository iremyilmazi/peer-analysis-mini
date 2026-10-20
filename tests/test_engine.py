import contextlib
import importlib
import io
import json
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from openpyxl.styles import Font

import banka_otomatik_guncelle as auto
from tests.helpers import workbook


class EngineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.input = self.root / 'input.xlsx'
        workbook(self.input)
        self.audit = self.root / 'audit.json'
        self.args = ['--workbook',str(self.input),'--history-dir',str(self.root/'history'),
                     '--cache-dir',str(self.root/'cache'),'--audit-report',str(self.audit),
                     '--backup-root',str(self.root/'backups')]

    def excel(self, spec, config, candidate, history):
        self.assertEqual('2026Q2',history.name)
        book = load_workbook(candidate)
        sheet = book[spec.company_name]
        for col, label in zip((77,78,79),('TL','YP','TOPLAM')):
            sheet.cell(1,col,'30.06.2026')
            sheet.cell(2,col,label)
            for row in range(3,21):
                sheet.cell(row,col,row*10)
        book.save(candidate)
        book.close()

    def run_case(self, *options, downloaded=None, bank='akbank', pdf_error=None):
        args = auto.build_parser(bank).parse_args(self.args+list(options))
        with patch.object(auto,'download_reports',return_value=downloaded or ['2026Q2']) as download, \
             patch.object(auto,'update_excel',side_effect=self.excel), \
             patch.object(auto,'update_pdf',side_effect=pdf_error), contextlib.redirect_stdout(io.StringIO()):
            result = auto.run_bank(auto.BANKS[bank],args)
        return result,json.loads(self.audit.read_text(encoding='utf-8')),download

    def local_sources(self, quarter='2026Q2', root=None, excel=None, pdf=None):
        # Signature-valid synthetic sources. Actual parsers are mocked in these
        # routing tests; no public financial data or real KAP request is used.
        folder = (root or self.root/'history')/quarter
        folder.mkdir(parents=True,exist_ok=True)
        (folder/'Bildirimler.xls').write_bytes(excel if excel is not None else b'<html>'+b' '*12000)
        (folder/'report.pdf').write_bytes(pdf if pdf is not None else b'%PDF-1.7\n'+b' '*12000)
        return folder

    def test_all_banks_reuse_local_target_without_kap_or_input_changes(self):
        folder = self.local_sources()
        before = {p:auto.sha256(p) for p in (self.input,*folder.iterdir())}
        with patch.object(auto,'kap_connection',side_effect=AssertionError('KAP must not open')) as connection:
            for key in auto.BANKS:
                with self.subTest(bank=key):
                    code, report, download = self.run_case('--dry-run',bank=key)
                    self.assertEqual((0,'OK'),(code,report['status']),report)
                    self.assertEqual(['30.06.2026'],report['new_periods'])
                    self.assertEqual([],report['downloaded_quarters'])
                    self.assertNotIn('1-KAP',[s['name'] for s in report['steps']])
                    download.assert_not_called()
            connection.assert_not_called()
        self.assertEqual(before,{p:auto.sha256(p) for p in before})

    def test_local_target_has_priority_over_legacy_force_download(self):
        self.local_sources()
        code, report, download = self.run_case('--dry-run','--force-download')
        self.assertEqual((0,'OK'),(code,report['status']))
        download.assert_not_called()

    def test_explicit_target_also_reuses_local_sources(self):
        self.local_sources()
        code, report, download = self.run_case('--dry-run','--period','30.06.2026')
        self.assertEqual((0,'OK'),(code,report['status']))
        self.assertEqual(['30.06.2026'],report['new_periods'])
        download.assert_not_called()

    def test_partial_or_empty_target_uses_downloader(self):
        for i,missing in enumerate(('Bildirimler.xls','report.pdf','both')):
            with self.subTest(missing=missing):
                root = self.root/f'partial-{i}'
                folder = self.local_sources(root=root)
                if missing=='both':
                    (folder/'Bildirimler.xls').unlink()
                    (folder/'report.pdf').unlink()
                else:
                    (folder/missing).unlink()
                code, report, download = self.run_case('--dry-run','--history-dir',str(root))
                self.assertEqual((0,'OK'),(code,report['status']))
                download.assert_called_once()
                self.assertEqual('31.03.2026',download.call_args.args[-1])

    def test_invalid_or_login_sources_do_not_bypass_downloader(self):
        cases = [(b'',None),(None,b''),(None,b'not-a-pdf'+b' '*12000),
                 (b'<html><form>login'+b' '*12000,None),(b'wrong-type'+b' '*12000,None)]
        for i,(excel,pdf) in enumerate(cases):
            with self.subTest(case=i):
                root = self.root/f'invalid-{i}'
                self.local_sources(root=root,excel=excel,pdf=pdf)
                code, report, download = self.run_case('--dry-run','--history-dir',str(root))
                self.assertEqual((0,'OK'),(code,report['status']))
                download.assert_called_once()

    def test_other_quarter_does_not_replace_expected_local_target(self):
        self.local_sources('2026Q1')
        self.local_sources('2026Q3')
        code, report, download = self.run_case('--dry-run')
        self.assertEqual((0,'OK'),(code,report['status']))
        self.assertEqual(['30.06.2026'],report['new_periods'])
        download.assert_called_once()

    def test_later_processing_failure_does_not_retry_kap_or_change_input(self):
        self.local_sources()
        before = auto.sha256(self.input)
        with patch.object(auto,'kap_connection',side_effect=AssertionError('KAP must not open')) as connection:
            code, report, download = self.run_case(pdf_error=AttributeError('missing bank function'))
        self.assertEqual((1,'ERROR'),(code,report['status']))
        self.assertEqual(before,auto.sha256(self.input))
        self.assertIsNone(report['backup'])
        self.assertEqual([],report['downloaded_quarters'])
        download.assert_not_called()
        connection.assert_not_called()

    def test_dry_run_keeps_input_and_creates_output(self):
        before = auto.sha256(self.input)
        code, report, _ = self.run_case('--dry-run')
        self.assertEqual((0,'OK'),(code,report['status']))
        self.assertEqual(before,auto.sha256(self.input))
        self.assertTrue(Path(report['output']).is_file())
        self.assertIsNone(report['backup'])

    def test_live_update_has_exact_backup(self):
        before = auto.sha256(self.input)
        code, report, _ = self.run_case()
        self.assertEqual((0,'OK'),(code,report['status']))
        self.assertEqual(before,auto.sha256(Path(report['backup'])))
        self.assertNotEqual(before,auto.sha256(self.input))
        self.assertIn('30.06.2026',auto.workbook_periods(self.input,'akbank'))

    def test_missing_existing_period_does_not_download_history(self):
        book = load_workbook(self.input)
        for col in (74,75,76): book['akbank'].cell(1,col).value = None
        book.save(self.input); book.close()
        before = auto.sha256(self.input)
        code, report, download = self.run_case()
        self.assertEqual((1,'ERROR'),(code,report['status']))
        download.assert_not_called()
        self.assertEqual(before,auto.sha256(self.input))

    def test_explicit_period_cannot_skip_a_quarter(self):
        code, report, download = self.run_case('--period','30.09.2026')
        self.assertEqual((1,'ERROR'),(code,report['status']))
        download.assert_not_called()

    def test_wrong_download_quarter_is_rejected(self):
        before = auto.sha256(self.input)
        code, report, _ = self.run_case(downloaded=['2026Q3'])
        self.assertEqual((1,'ERROR'),(code,report['status']))
        self.assertEqual(before,auto.sha256(self.input))

    def test_skip_download_does_not_select_q3(self):
        with patch.object(auto,'complete_history_periods',return_value={'30.09.2026'}):
            code, report, download = self.run_case('--skip-download')
        self.assertEqual((0,'NO_NEW_PERIOD'),(code,report['status']))
        download.assert_not_called()

    def test_dry_run_output_cannot_be_input(self):
        before = auto.sha256(self.input)
        code, report, _ = self.run_case('--dry-run','--dry-run-output',str(self.input))
        self.assertEqual((1,'ERROR'),(code,report['status']))
        self.assertEqual(before,auto.sha256(self.input))

    def test_external_input_change_prevents_commit(self):
        original_audit = auto.audit_candidate
        def external_change(*args):
            outcome = original_audit(*args)
            book = load_workbook(self.input)
            book['akbank']['BV3'] = 666
            book.save(self.input); book.close()
            return outcome
        with patch.object(auto,'audit_candidate',side_effect=external_change):
            code, report, _ = self.run_case()
        self.assertEqual((1,'ERROR'),(code,report['status']))
        self.assertIsNone(report['backup'])
        book = load_workbook(self.input)
        self.assertEqual(666,book['akbank']['BV3'].value)
        self.assertIsNone(book['akbank']['BY1'].value)
        book.close()

    def test_old_period_value_or_style_change_is_rejected(self):
        candidate = self.root / 'candidate.xlsx'
        for style_change in (False,True):
            workbook(candidate)
            book = load_workbook(candidate)
            if style_change: book['akbank']['BV3'].font = Font(bold=True)
            else: book['akbank']['BV3'] = 999
            book.save(candidate); book.close()
            with self.assertRaisesRegex(RuntimeError,'History/veri deseni'):
                auto.audit_candidate(self.input,candidate,'akbank',set())

    def test_quarter_rollover(self):
        self.assertEqual('31.03.2027',auto.next_period('31.12.2026'))
        self.assertEqual('30.06.2026',auto.next_period('31.03.2026'))

    def test_all_banks_have_runtime_dependencies(self):
        import panel
        self.assertEqual(13,len(auto.BANKS))
        self.assertEqual(set(auto.BANKS),{b.key for b in panel.BANKS})
        for bank in panel.BANKS: self.assertTrue(bank.script.is_file(),bank.key)
        for spec in auto.BANKS.values():
            self.assertTrue(callable(importlib.import_module(spec.downloader_module).run_incremental))
            self.assertTrue(callable(getattr(importlib.import_module(spec.pdf_module),spec.pdf_function)))
        from ing_bank.ing_guncelle import finalize_ing_sheet
        self.assertTrue(callable(finalize_ing_sheet))

    def test_all_downloaders_wait_for_exact_next_quarter(self):
        for key,spec in auto.BANKS.items():
            module = importlib.import_module(spec.downloader_module)
            base = getattr(module,'_base',module)
            with self.subTest(bank=key), \
                 patch.object(base,'iter_candidate_disclosures',return_value=[]), \
                 patch.object(base,'choose_nc_by_quarter',return_value={'2026Q3':object()}), \
                 patch.object(base,'download_disclosure') as download, contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual([],auto.download_reports(spec,self.root/'empty',date(2026,10,19),False,'31.03.2026'))
                download.assert_not_called()
