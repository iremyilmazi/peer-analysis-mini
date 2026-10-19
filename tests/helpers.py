from pathlib import Path
from openpyxl import Workbook
import banka_otomatik_guncelle as auto


def workbook(path: Path):
    book = Workbook()
    book.remove(book.active)
    for spec in auto.BANKS.values():
        sheet = book.create_sheet(spec.company_name)
        sheet['A1'] = 'Kalem'
        sheet['A2'] = 'Açıklama'
        for col, label in zip((74,75,76), ('TL','YP','TOPLAM')):
            sheet.cell(1,col,'31.03.2026')
            sheet.cell(2,col,label)
        for row in range(3,21):
            sheet.cell(row,1,f'Finansal satır {row}')
            sheet.cell(row,74,row)
            sheet.cell(row,75,row*2)
            sheet.cell(row,76,row*3)
    book.save(path)
    book.close()
